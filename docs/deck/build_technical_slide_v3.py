#!/usr/bin/env python3
"""SIH 2026 deck, slide 3 (TECHNICAL APPROACH), v3: same architecture as v2 in a restrained engineering-figure style.
Palette: navy + greyscale; teal only for sensor inputs; lavender only for designed / next-build. No shadows, no pills.
Layout: INPUT (left) -> EMBEDDED CONTROL (centre) -> PHYSICAL OUTPUT (right), TECHNOLOGIES TO BE USED band, minimal flow.
Only the team's own photographs; numbers from the repository's recorded evidence.
    python3 docs/deck/build_technical_slide_v3.py   -> docs/deck/AquaSDR_slide3_technical_approach_v3.pptx
"""
import pathlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn

HERE = pathlib.Path(__file__).parent
IMG = HERE / "img"
OUT = HERE / "AquaSDR_slide3_technical_approach_v3.pptx"
rgb = RGBColor.from_string
NAVY, TEAL, BG, LAV, CHAR, GRAY, WHITE = rgb("183153"), rgb("3B8C7A"), rgb("F3F6F9"), rgb("8A7BAF"), rgb("263238"), rgb("667085"), rgb("FFFFFF")
RULE = rgb("D0D5DD")
FONT = "Arial"

prs = Presentation(HERE / "template" / "SIH2026-IDEA-template.pptx")
ids = prs.slides._sldIdLst
for i, sid in reversed(list(enumerate(list(ids)))):
    if i != 2: prs.part.drop_rel(sid.rId); ids.remove(sid)
s = prs.slides[0]
for sh in list(s.shapes):
    if sh.has_text_frame and "Technologies to be used" in sh.text_frame.text:
        sh._element.getparent().remove(sh._element)
    elif sh.has_text_frame and sh.text_frame.text.strip() == "TECHNICAL APPROACH":
        p = sh.text_frame.paragraphs[0]
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
        p.runs[0].text = "AquaSDR · TECHNICAL APPROACH"; p.runs[0].font.size = Pt(30)

I = Inches
def flat(shape):
    el = shape._element
    st = el.find(qn("p:style"))
    if st is not None: el.remove(st)
    spPr = el.find(qn("p:spPr"))
    if spPr is not None and spPr.find(qn("a:effectLst")) is None: spPr.append(spPr.makeelement(qn("a:effectLst"), {}))
    return shape

def text(x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, size=8.5, color=CHAR, bold=False, italic=False):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, side, 0)
    if isinstance(paras, str): paras = [paras]
    for k, para in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); p.alignment = align
        for t, st in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run(); r.text = t; f = r.font; f.name = FONT
            f.size = Pt(st.get("size", size)); f.bold = st.get("bold", bold); f.italic = st.get("italic", italic); f.color.rgb = st.get("color", color)
    return tb

def rect(x, y, w, h, fill=None, line=NAVY, lw=0.75, dashed=False, radius=0.0):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
    if radius: b.adjustments[0] = radius
    if fill is None: b.fill.background()
    else: b.fill.solid(); b.fill.fore_color.rgb = fill
    if line is None: b.line.fill.background()
    else:
        b.line.color.rgb = line; b.line.width = Pt(lw)
        if dashed: b.line.dash_style = MSO_LINE.DASH
    return flat(b)

def seg(x1, y1, x2, y2, color=NAVY, lw=1.0, head=False, dashed=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = color; c.line.width = Pt(lw); flat(c)
    if dashed: c.line.dash_style = MSO_LINE.DASH
    if head:
        ln = c.line._get_or_add_ln(); ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "sm", "len": "sm"}))
    return c

def path(pts, color=NAVY, lw=1.0):
    for k in range(len(pts) - 1): seg(*pts[k], *pts[k + 1], color=color, lw=lw, head=(k == len(pts) - 2))

def dot(x, y, r=0.035, line=NAVY):
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, I(x - r), I(y - r), I(2 * r), I(2 * r))
    d.fill.solid(); d.fill.fore_color.rgb = WHITE; d.line.color.rgb = line; d.line.width = Pt(1.0); flat(d)

def photo(file, x, y, w, h, line=NAVY, lw=0.75):
    p = s.shapes.add_picture(str(file), I(x), I(y), I(w), I(h)); p.line.color.rgb = line; p.line.width = Pt(lw)
    return flat(p)

def header(x, y, w, label, subtitle):
    text(x, y, w, 0.17, [[(label, {"bold": True, "size": 9, "color": NAVY})]])
    text(x, y + 0.17, w, 0.17, [[(subtitle, {"size": 8.2, "color": GRAY})]])
    seg(x, y + 0.38, x + w, y + 0.38, color=RULE, lw=0.75)

text(2.0, 0.93, 9.3, 0.24, [[("Adaptive Software-Defined Sonar Transmitter for AUVs", {"size": 12, "italic": True, "color": GRAY})]], align=PP_ALIGN.CENTER)

# =========================== INPUT ===========================
header(0.30, 1.30, 3.25, "SENSE", "Real environment · real sensors")
PX, PY, PW = 1.55, 1.78, 1.85
PH = PW / 0.909
photo(IMG / "sensors.jpg", PX, PY, PW, PH)
P = lambda u, v: (PX + u / 1400 * PW, PY + v / 1541 * PH)
for label, (u, v), ly in [("TDS / SALINITY PROXY", (514, 424), 2.00), ("HC-SR04 · AIR RANGE", (719, 617), 2.36),
                          ("TURBIDITY", (540, 822), 2.72), ("DS18B20 TEMPERATURE", (1079, 796), 3.08), ("PRESSURE / DEPTH", (1028, 1130), 3.44)]:
    px, py = P(u, v)
    text(0.20, ly - 0.08, 1.30, 0.17, [[(label, {"bold": True, "size": 6.9, "color": NAVY})]], align=PP_ALIGN.RIGHT)
    seg(1.53, ly, px, py, color=NAVY, lw=0.6); dot(px, py)
text(PX, PY + PH + 0.03, PW, 0.17, [[("ADC · 1-Wire · I²C", {"bold": True, "size": 7.8, "color": TEAL})]], align=PP_ALIGN.CENTER)
text(0.30, 4.06, 3.25, 0.32, [[("Also on board  ", {"bold": True, "size": 7.2, "color": NAVY}), ("BMP390 · IMU + magnetometer (I²C) · INA219 ×2 mounted, not commissioned", {"size": 7.2, "color": GRAY})]])
text(0.30, 4.50, 3.25, 0.17, [[("ONE POLICY · TWO INPUTS", {"bold": True, "size": 7.4, "color": NAVY})]])
for k, (a, b) in enumerate([("A0 sensor", "TDS_PROXY → ADC"), ("Web command", "console → UART")]):
    bx = 0.30 + k * 1.55
    rect(bx, 4.70, 1.45, 0.40, fill=WHITE, line=TEAL, lw=0.75)
    text(bx, 4.73, 1.45, 0.36, [[(a, {"bold": True, "size": 7.6, "color": NAVY})], [(b, {"size": 6.9, "color": GRAY})]], align=PP_ALIGN.CENTER)
for bx in (1.025, 2.575): seg(bx, 5.10, bx, 5.24, color=TEAL, lw=0.9)
seg(1.025, 5.24, 3.66, 5.24, color=TEAL, lw=0.9)

# =========================== EMBEDDED CONTROL ===========================
header(3.80, 1.30, 4.60, "DECIDE · SYNTHESISE · STREAM", "STM32F446RE embedded control · methodology & process")
rect(3.80, 1.76, 4.60, 3.98, fill=BG, line=NAVY, lw=0.75)
rect(3.80, 1.76, 4.60, 0.27, fill=NAVY, line=None)
text(3.80, 1.76, 4.60, 0.27, [[("THE NEXT PING IS COMPUTED ON THE MCU", {"bold": True, "size": 9.5, "color": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
BX, BY, BW = 3.90, 2.11, 3.02
BH = BW / 2.7027
photo(IMG / "bay.jpg", BX, BY, BW, BH)
for label, (fx, fy), (lx, ly) in [("STM32 NUCLEO-F446RE", (0.278, 0.114), (0.03, 0.72)), ("Sensor interface boards", (0.683, 0.82), (0.57, 0.07))]:
    tx, ty = BX + fx * BW, BY + fy * BH
    w = 0.056 * len(label) + 0.12; lx_, ly_ = BX + lx * BW, BY + ly * BH
    seg(lx_ + w / 2, ly_ + (0 if ly_ > ty else 0.16), tx, ty, color=WHITE, lw=0.9); dot(tx, ty)
    rect(lx_, ly_, w, 0.16, fill=WHITE, line=NAVY, lw=0.5)
    text(lx_, ly_ + 0.005, w, 0.15, [[(label, {"bold": True, "size": 6.4, "color": NAVY})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rect(7.00, BY, 1.30, BH, fill=WHITE, line=RULE, lw=0.75)
text(7.07, BY + 0.03, 1.20, BH - 0.06, [[("MCU ADC CAPTURE", {"bold": True, "size": 7.4, "color": NAVY})],
                                        [("LFM example · on-board", {"size": 6.6, "italic": True, "color": GRAY})],
                                        [("Centre ", {"size": 7, "color": GRAY}), ("42 kHz", {"bold": True, "size": 7})],
                                        [("Sweep ", {"size": 7, "color": GRAY}), ("4 kHz", {"bold": True, "size": 7})],
                                        [("Pulse ", {"size": 7, "color": GRAY}), ("2 ms", {"bold": True, "size": 7}), ("  Hann", {"size": 7, "color": GRAY})],
                                        [("330/330", {"bold": True, "size": 7}), (" cmds ACKed", {"size": 6.8, "color": GRAY})],
                                        [("0", {"bold": True, "size": 7}), (" DMA underruns", {"size": 6.8, "color": GRAY})]], anchor=MSO_ANCHOR.MIDDLE)
steps = [("01", "ACQUIRE", "ADC · 1-Wire · I²C · stale / out-of-range inputs rejected"),
         ("02", "DECIDE", "smoothing filter + hysteresis → CLEAR / TRANSITION / MURKY"),
         ("03", "SYNTHESISE", "LFM · geometric sweep · Barker-13 × Hann / Hamming / Blackman"),
         ("04", "STREAM", "TIM6 + DMA → 12-bit DAC1 / PA4 · 1 MS/s · 0 underruns"),
         ("06", "SELF-VALIDATE", "ADC + DMA capture of the DAC output → FFT · spectrogram · matched filter")]
Y0, DY, NX = 3.37, 0.46, 3.95
seg(NX + 0.11, Y0 + 0.21, NX + 0.11, Y0 + 4 * DY - 0.03, color=RULE, lw=1.0)
for k, (n, t, d) in enumerate(steps):
    y = Y0 + k * DY
    rect(NX - 0.02, y - 0.01, 0.26, 0.2, fill=BG, line=None)
    text(NX, y, 0.24, 0.18, [[(n, {"bold": True, "size": 8.4, "color": NAVY})]], align=PP_ALIGN.CENTER)
    text(4.30, y, 4.05, 0.18, [[(t, {"bold": True, "size": 8.8, "color": NAVY})]])
    text(4.30, y + 0.18, 4.05, 0.2, [[(d, {"size": 7.3, "color": GRAY})]])
path([(3.86, Y0 + 3 * DY + 0.09), (3.86, Y0 + 4 * DY + 0.09), (NX - 0.02, Y0 + 4 * DY + 0.09)], color=NAVY, lw=0.75)   # DAC pin -> capture
seg(PX + PW, Y0 + 0.09, 3.80, Y0 + 0.09, color=TEAL, lw=1.1, head=True)                                       # sensors -> 01
path([(3.66, 5.24), (3.66, Y0 + DY + 0.09), (3.80, Y0 + DY + 0.09)], color=TEAL, lw=1.1)                        # inputs -> 02

# =========================== PHYSICAL OUTPUT ===========================
header(8.65, 1.30, 2.95, "NEXT PING", "Adaptive waveform parameters")
LX, CW, TY = 8.65, 0.70, 1.80
text(LX, TY, 0.86, 0.18, "")
for j, c in enumerate(["CLEAR", "TRANSITION", "MURKY"]):
    text(LX + 0.86 + j * CW, TY, CW, 0.18, [[(c, {"bold": True, "size": 7.2, "color": NAVY})]], align=PP_ALIGN.CENTER)
seg(LX, TY + 0.2, LX + 0.86 + 3 * CW, TY + 0.2, color=NAVY, lw=0.75)
for i, (n, vals) in enumerate([("Centre", ["42 kHz", "40 kHz", "38 kHz"]), ("Bandwidth", ["4 kHz", "2 kHz", "1 kHz"]),
                               ("Pulse", ["2 ms", "4 ms", "8 ms"]), ("Amplitude", ["35 %", "50 %", "62 %"])]):
    y = TY + 0.25 + i * 0.2
    text(LX, y, 0.86, 0.18, [[(n.upper(), {"bold": True, "size": 7.2, "color": NAVY})]])
    for j, v in enumerate(vals): text(LX + 0.86 + j * CW, y, CW, 0.18, [[(v, {"size": 8.2, "color": CHAR})]], align=PP_ALIGN.CENTER)
    seg(LX, y + 0.19, LX + 0.86 + 3 * CW, y + 0.19, color=RULE, lw=0.5)
text(LX, 2.91, 4.45, 0.17, [[("Water condition changes → next-ping parameters change", {"size": 7.4, "italic": True, "color": GRAY})]])
path([(8.40, Y0 + DY + 0.09), (8.53, Y0 + DY + 0.09), (8.53, 2.40), (8.62, 2.40)], color=NAVY, lw=1.0)            # 02 -> next ping

rect(8.65, 3.14, 4.45, 1.02, fill=WHITE, line=LAV, lw=0.9, dashed=True)
text(8.73, 3.19, 3.5, 0.18, [[("05  CONDITION · ANALOG FRONT END", {"bold": True, "size": 8.4, "color": NAVY})]])
photo(IMG / "afe_chain.png", 8.73, 3.43, 2.40, 2.40 / 3.645, line=WHITE, lw=0.25)
text(11.25, 3.43, 1.8, 0.5, [[("2nd-order low-pass filter", {"bold": True, "size": 7.2, "color": NAVY})], [("70.7 kHz design target", {"size": 7, "color": GRAY})],
                             [("TLV9062 ×1.5", {"bold": True, "size": 7.2, "color": NAVY})]])
text(11.25, 3.93, 1.8, 0.16, [[("DESIGNED · NEXT BUILD", {"bold": True, "size": 6.8, "color": LAV})]])
path([(8.40, Y0 + 3 * DY + 0.09), (8.57, Y0 + 3 * DY + 0.09), (8.57, 3.98), (8.65, 3.98)], color=NAVY, lw=1.0)    # 04 DAC -> 05

seg(10.875, 4.16, 10.875, 4.27, color=NAVY, lw=1.0, head=True)
PW2 = 4.45; PH2 = PW2 / 3.735
photo(IMG / "payload.jpg", 8.65, 4.28, PW2, PH2, lw=1.0)
text(8.65, 4.28 + PH2 + 0.03, PW2, 0.17, [[("AQUASDR PAYLOAD · PHYSICAL PROTOTYPE", {"bold": True, "size": 8, "color": NAVY}), ("   acrylic payload hull", {"size": 7.2, "color": GRAY})]])
text(8.65, 4.28 + PH2 + 0.21, PW2, 0.16, [[("External scope validation → PENDING  ·  AUV hull-slot pod → NEXT BUILD", {"size": 6.9, "italic": True, "color": GRAY})]])

# =========================== TECHNOLOGIES TO BE USED ===========================
TY2 = 5.86
rect(0.30, TY2, 12.80, 0.62, fill=WHITE, line=RULE, lw=0.75)
text(0.40, TY2 + 0.12, 1.25, 0.4, [[("TECHNOLOGIES", {"bold": True, "size": 8, "color": NAVY})], [("TO BE USED", {"bold": True, "size": 8, "color": NAVY})]])
seg(1.62, TY2 + 0.08, 1.62, TY2 + 0.54, color=RULE, lw=0.75)
techs = [("Programming", "C / C++ · Python · JavaScript"), ("Embedded", "STM32F446RE · ADC · TIM6 · DMA · 12-bit DAC"),
         ("Signal processing", "LFM · geometric sweep · Barker-13 · FFT · windowing"), ("Analog", "Sallen-Key LPF · TLV9062 op-amp"),
         ("Software", "React · Node.js / Express · WebSocket · OpenSCAD"), ("Hardware", "Turbidity · TDS · pressure · DS18B20 · IMU · BMP390 · INA219")]
for k, (a, b) in enumerate(techs):
    cx, cy = 1.74 + (k % 3) * 3.79, TY2 + 0.09 + (k // 3) * 0.24
    text(cx, cy, 3.75, 0.2, [[(a + "  ", {"bold": True, "size": 7, "color": NAVY}), (b, {"size": 7, "color": CHAR})]])

# =========================== flow + legend ===========================
flow = ["SENSE", "DECIDE", "SYNTHESISE", "STREAM", "CONDITION", "VERIFY"]
widths = [0.78, 0.82, 1.08, 0.86, 1.04, 0.82]; gap = 0.42; x = 0.55
for k, (t, w) in enumerate(zip(flow, widths)):
    rect(x, 6.60, w, 0.24, fill=WHITE, line=LAV if t == "CONDITION" else NAVY, lw=0.75, dashed=(t == "CONDITION"))
    text(x, 6.60, w, 0.24, [[(t, {"bold": True, "size": 7.8, "color": NAVY})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if k < len(flow) - 1: seg(x + w + 0.06, 6.72, x + w + gap - 0.06, 6.72, color=NAVY, lw=0.9, head=True)
    x += w + gap
seg(9.70, 6.72, 10.05, 6.72, color=NAVY, lw=1.0); text(10.12, 6.64, 1.2, 0.16, [[("built / running", {"size": 7, "color": GRAY})]])
seg(11.25, 6.72, 11.60, 6.72, color=LAV, lw=1.0, dashed=True); text(11.67, 6.64, 1.4, 0.16, [[("designed / next build", {"size": 7, "color": GRAY})]])

prs.save(OUT)
print("wrote", OUT)
