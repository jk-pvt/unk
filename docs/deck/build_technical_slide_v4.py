#!/usr/bin/env python3
"""SIH 2026 deck, slide 3 (TECHNICAL APPROACH), v4 (readability pass on v3): same architecture as v2 in a restrained engineering-figure style.
Palette: navy + greyscale; teal only for sensor inputs; lavender only for designed / next-build. No shadows, no pills.
Layout: INPUT (left) -> EMBEDDED CONTROL (centre) -> PHYSICAL OUTPUT (right), TECHNOLOGIES TO BE USED band, minimal flow.
Only the team's own photographs; numbers from the repository's recorded evidence.
    python3 docs/deck/build_technical_slide_v3.py   -> docs/deck/AquaSDR_slide3_technical_approach_v4.pptx
"""
import pathlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn

HERE = pathlib.Path(__file__).parent
IMG = HERE / "img"
OUT = HERE / "AquaSDR_slide3_technical_approach_v4.pptx"
rgb = RGBColor.from_string
NAVY, TEAL, BG, LAV, CHAR, GRAY, WHITE = rgb("183153"), rgb("3B8C7A"), rgb("F3F6F9"), rgb("8A7BAF"), rgb("263238"), rgb("667085"), rgb("FFFFFF")
RULE = rgb("D0D5DD")
FONT = "Arial"

prs = Presentation(HERE / "template" / "SIH2026-IDEA-template.pptx")
ids = prs.slides._sldIdLst
for i, sid in reversed(list(enumerate(list(ids)))):
    if i != 2: prs.part.drop_rel(sid.rId); ids.remove(sid)
s = prs.slides[0]
for sh in list(s.shapes):
    if sh.has_text_frame and "Technologies to be used" in sh.text_frame.text:
        sh._element.getparent().remove(sh._element)
    elif sh.has_text_frame and sh.text_frame.text.strip() == "TECHNICAL APPROACH":
        p = sh.text_frame.paragraphs[0]
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
        p.runs[0].text = "AquaSDR · TECHNICAL APPROACH"; p.runs[0].font.size = Pt(30)

I = Inches
def flat(shape):
    el = shape._element
    st = el.find(qn("p:style"))
    if st is not None: el.remove(st)
    spPr = el.find(qn("p:spPr"))
    if spPr is not None and spPr.find(qn("a:effectLst")) is None: spPr.append(spPr.makeelement(qn("a:effectLst"), {}))
    return shape

def text(x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, size=8.5, color=CHAR, bold=False, italic=False):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, side, 0)
    if isinstance(paras, str): paras = [paras]
    for k, para in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); p.alignment = align
        for t, st in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run(); r.text = t; f = r.font; f.name = FONT
            f.size = Pt(st.get("size", size)); f.bold = st.get("bold", bold); f.italic = st.get("italic", italic); f.color.rgb = st.get("color", color)
    return tb

def rect(x, y, w, h, fill=None, line=NAVY, lw=0.75, dashed=False, radius=0.0):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
    if radius: b.adjustments[0] = radius
    if fill is None: b.fill.background()
    else: b.fill.solid(); b.fill.fore_color.rgb = fill
    if line is None: b.line.fill.background()
    else:
        b.line.color.rgb = line; b.line.width = Pt(lw)
        if dashed: b.line.dash_style = MSO_LINE.DASH
    return flat(b)

def seg(x1, y1, x2, y2, color=NAVY, lw=1.0, head=False, dashed=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = color; c.line.width = Pt(lw); flat(c)
    if dashed: c.line.dash_style = MSO_LINE.DASH
    if head:
        ln = c.line._get_or_add_ln(); ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "sm", "len": "sm"}))
    return c

def path(pts, color=NAVY, lw=1.0):
    for k in range(len(pts) - 1): seg(*pts[k], *pts[k + 1], color=color, lw=lw, head=(k == len(pts) - 2))

def dot(x, y, r=0.035, line=NAVY):
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, I(x - r), I(y - r), I(2 * r), I(2 * r))
    d.fill.solid(); d.fill.fore_color.rgb = WHITE; d.line.color.rgb = line; d.line.width = Pt(1.0); flat(d)

def photo(file, x, y, w, h, line=NAVY, lw=0.75):
    p = s.shapes.add_picture(str(file), I(x), I(y), I(w), I(h)); p.line.color.rgb = line; p.line.width = Pt(lw)
    return flat(p)

def tag(x, y, label="BUILT"):
    w = 0.075 * len(label) + 0.16
    rect(x, y, w, 0.17, fill=NAVY, line=None)
    text(x, y, w, 0.17, [[(label, {"bold": True, "size": 6.8, "color": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def header(x, y, w, label, subtitle):
    text(x, y, w, 0.18, [[(label, {"bold": True, "size": 9.6, "color": NAVY})]])
    text(x, y + 0.18, w, 0.17, [[(subtitle, {"size": 8.6, "color": GRAY})]])
    seg(x, y + 0.38, x + w, y + 0.38, color=RULE, lw=0.75)

text(2.0, 0.93, 9.3, 0.24, [[("Adaptive Software-Defined Sonar Transmitter for AUVs", {"size": 12, "italic": True, "color": GRAY})]], align=PP_ALIGN.CENTER)

# =========================== INPUT ===========================
header(0.30, 1.30, 3.25, "SENSE", "Real environment · real sensors")
PX, PY, PW = 1.55, 1.78, 1.85
PH = PW / 0.909
photo(IMG / "sensors.jpg", PX, PY, PW, PH)
tag(PX + PW - 0.58, PY + 0.05)
P = lambda u, v: (PX + u / 1400 * PW, PY + v / 1541 * PH)
for label, (u, v), ly in [("TDS / SALINITY PROXY", (514, 424), 2.00), ("HC-SR04 · AIR RANGE", (719, 617), 2.36),
                          ("TURBIDITY", (540, 822), 2.72), ("DS18B20 TEMPERATURE", (1079, 796), 3.08), ("PRESSURE / DEPTH", (1028, 1130), 3.44)]:
    px, py = P(u, v)
    text(0.05, ly - 0.09, 1.45, 0.19, [[(label, {"bold": True, "size": 7.6, "color": NAVY})]], align=PP_ALIGN.RIGHT)
    ex = PX; ey = ly + (ex - 1.53) / (px - 1.53) * (py - ly)          # where the leader enters the photo
    seg(1.53, ly, ex, ey, color=NAVY, lw=0.9)
    seg(ex, ey, px, py, color=WHITE, lw=1.4); dot(px, py, r=0.042)
text(PX, PY + PH + 0.03, PW, 0.17, [[("ADC · 1-Wire · I²C", {"bold": True, "size": 8.6, "color": TEAL})]], align=PP_ALIGN.CENTER)
text(0.30, 4.06, 3.25, 0.32, [[("Also on board  ", {"bold": True, "size": 7.8, "color": NAVY}), ("BMP390 · IMU + magnetometer · INA219 ×2 (mounted, not commissioned)", {"size": 7.8, "color": GRAY})]])
text(0.30, 4.50, 3.25, 0.17, [[("ONE POLICY · TWO INPUTS", {"bold": True, "size": 8.2, "color": NAVY})]])
for k, (a, b) in enumerate([("A0 sensor", "TDS_PROXY → ADC"), ("Web command", "console → UART")]):
    bx = 0.30 + k * 1.55
    rect(bx, 4.70, 1.45, 0.40, fill=WHITE, line=TEAL, lw=0.75)
    text(bx, 4.73, 1.45, 0.36, [[(a, {"bold": True, "size": 8.4, "color": NAVY})], [(b, {"size": 7.6, "color": GRAY})]], align=PP_ALIGN.CENTER)
for bx in (1.025, 2.575): seg(bx, 5.10, bx, 5.24, color=TEAL, lw=0.9)
seg(1.025, 5.24, 3.66, 5.24, color=TEAL, lw=0.9)

# =========================== EMBEDDED CONTROL ===========================
header(3.80, 1.30, 4.60, "DECIDE · SYNTHESISE · STREAM", "STM32F446RE embedded control")
rect(3.80, 1.76, 4.60, 3.98, fill=BG, line=NAVY, lw=0.75)
rect(3.80, 1.76, 4.60, 0.27, fill=NAVY, line=None)
text(3.80, 1.76, 4.60, 0.27, [[("THE NEXT PING IS COMPUTED ON THE MCU", {"bold": True, "size": 10.5, "color": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
BX, BY, BW = 3.90, 2.10, 2.88
BH = BW / 2.7027
photo(IMG / "bay.jpg", BX, BY, BW, BH)
tag(BX + 0.05, BY + 0.05)
for label, (fx, fy), (lx, ly) in [("STM32 NUCLEO-F446RE", (0.278, 0.114), (0.03, 0.72)), ("Sensor interface boards", (0.683, 0.82), (0.40, 0.07))]:
    tx, ty = BX + fx * BW, BY + fy * BH
    w = 0.062 * len(label) + 0.14; lx_, ly_ = BX + lx * BW, BY + ly * BH
    seg(lx_ + w / 2, ly_ + (0 if ly_ > ty else 0.16), tx, ty, color=WHITE, lw=0.9); dot(tx, ty)
    rect(lx_, ly_, w, 0.18, fill=WHITE, line=NAVY, lw=0.5)
    text(lx_, ly_ + 0.01, w, 0.16, [[(label, {"bold": True, "size": 7.0, "color": NAVY})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rect(6.86, BY, 1.44, BH, fill=WHITE, line=NAVY, lw=0.75)
text(6.93, BY + 0.02, 1.33, BH - 0.04, [[("MCU ADC CAPTURE", {"bold": True, "size": 7.6, "color": NAVY}), (" · LFM", {"size": 7.2, "color": GRAY})],
                                        [("42 kHz", {"bold": True, "size": 11.5, "color": NAVY})],
                                        [("4 kHz BW", {"bold": True, "size": 11.5, "color": NAVY})],
                                        [("2 ms", {"bold": True, "size": 11.5, "color": NAVY})],
                                        [("Hann · 330/330 commands", {"size": 6.8, "color": GRAY})],
                                        [("0 DMA underruns", {"size": 6.8, "color": GRAY})]], anchor=MSO_ANCHOR.MIDDLE)
steps = [("01", "ACQUIRE", "ADC · 1-Wire · I²C", "Reject stale / invalid inputs"),
         ("02", "DECIDE", "Smoothing + hysteresis", "CLEAR / TRANSITION / MURKY"),
         ("03", "SYNTHESISE", "LFM · Geometric · Barker-13", "Hann · Hamming · Blackman"),
         ("04", "STREAM", "TIM6 + DMA → 12-bit DAC1 / PA4", "1 MS/s · 0 underruns · CPU free between transfers"),
         ("06", "SELF-VALIDATE", "ADC + DMA capture of DAC output", "FFT · spectrogram · matched filter")]
Y0, DY, NX = 3.22, 0.495, 3.95
seg(NX + 0.12, Y0 + 0.22, NX + 0.12, Y0 + 4 * DY - 0.03, color=RULE, lw=1.0)
for k, (n, t, d1, d2) in enumerate(steps):
    y = Y0 + k * DY
    rect(NX - 0.02, y - 0.01, 0.28, 0.21, fill=BG, line=None)
    text(NX, y, 0.26, 0.2, [[(n, {"bold": True, "size": 9.2, "color": NAVY})]], align=PP_ALIGN.CENTER)
    text(4.32, y - 0.01, 4.05, 0.2, [[(t, {"bold": True, "size": 9.6, "color": NAVY})]])
    text(4.32, y + 0.19, 4.05, 0.3, [[(d1, {"size": 8.0, "color": CHAR})], [(d2, {"size": 8.0, "color": GRAY})]])
path([(3.86, Y0 + 3 * DY + 0.09), (3.86, Y0 + 4 * DY + 0.09), (NX - 0.02, Y0 + 4 * DY + 0.09)], color=NAVY, lw=0.75)   # DAC pin -> capture
seg(PX + PW, Y0 + 0.09, 3.80, Y0 + 0.09, color=TEAL, lw=1.1, head=True)                                       # sensors -> 01
path([(3.66, 5.24), (3.66, Y0 + DY + 0.09), (3.80, Y0 + DY + 0.09)], color=TEAL, lw=1.1)                        # inputs -> 02

# =========================== PHYSICAL OUTPUT ===========================
header(8.65, 1.30, 2.95, "ADAPTED NEXT PING", "Adaptive waveform parameters")
LX, CW, TY = 8.65, 0.70, 1.80
text(LX, TY, 0.86, 0.18, "")
for j, c in enumerate(["CLEAR", "TRANSITION", "MURKY"]):
    text(LX + 0.86 + j * CW, TY, CW, 0.18, [[(c, {"bold": True, "size": 7.8, "color": NAVY})]], align=PP_ALIGN.CENTER)
seg(LX, TY + 0.2, LX + 0.86 + 3 * CW, TY + 0.2, color=NAVY, lw=0.75)
for i, (n, vals) in enumerate([("Centre", ["42 kHz", "40 kHz", "38 kHz"]), ("Bandwidth", ["4 kHz", "2 kHz", "1 kHz"]),
                               ("Pulse", ["2 ms", "4 ms", "8 ms"]), ("Amplitude", ["35 %", "50 %", "62 %"])]):
    y = TY + 0.25 + i * 0.2
    text(LX, y, 0.86, 0.18, [[(n.upper(), {"bold": True, "size": 7.8, "color": NAVY})]])
    for j, v in enumerate(vals): text(LX + 0.86 + j * CW, y, CW, 0.18, [[(v, {"size": 8.8, "color": CHAR})]], align=PP_ALIGN.CENTER)
    seg(LX, y + 0.19, LX + 0.86 + 3 * CW, y + 0.19, color=RULE, lw=0.5)
text(LX, 2.90, 4.45, 0.34, [[("CLEAR → shorter / sharper ping", {"size": 8.0, "color": CHAR})], [("MURKY → lower band / longer pulse / higher output demand", {"size": 8.0, "color": CHAR})]])
path([(8.40, Y0 + DY + 0.09), (8.53, Y0 + DY + 0.09), (8.53, 2.40), (8.62, 2.40)], color=NAVY, lw=1.0)            # 02 -> next ping

rect(8.65, 3.30, 4.45, 0.98, fill=WHITE, line=LAV, lw=0.9, dashed=True)
text(8.73, 3.35, 3.5, 0.18, [[("05  CONDITION · ANALOG FRONT END", {"bold": True, "size": 9.0, "color": NAVY})]])
photo(IMG / "afe_chain.png", 8.73, 3.60, 2.30, 2.30 / 3.645, line=WHITE, lw=0.25)
text(11.13, 3.58, 1.95, 0.5, [[("2nd-order low-pass filter", {"bold": True, "size": 7.8, "color": NAVY})], [("70.7 kHz design target", {"size": 7.6, "color": GRAY})],
                             [("TLV9062 ×1.5", {"bold": True, "size": 7.8, "color": NAVY})]])
text(11.13, 4.05, 1.95, 0.17, [[("DESIGNED · BUILDING NEXT", {"bold": True, "size": 7.4, "color": LAV})]])
path([(8.40, Y0 + 3 * DY + 0.09), (8.57, Y0 + 3 * DY + 0.09), (8.57, 4.15), (8.65, 4.15)], color=NAVY, lw=1.0)    # 04 DAC -> 05

seg(10.875, 4.28, 10.875, 4.37, color=NAVY, lw=1.0, head=True)
PW2 = 4.20; PH2 = PW2 / 3.735
photo(IMG / "payload.jpg", 8.775, 4.38, PW2, PH2, lw=1.0)
tag(8.775 + 0.06, 4.38 + 0.06)
text(8.65, 4.38 + PH2 + 0.03, 4.45, 0.18, [[("AQUASDR PAYLOAD · PHYSICAL PROTOTYPE", {"bold": True, "size": 8.6, "color": NAVY}), ("  acrylic hull", {"size": 7.8, "color": GRAY})]])
text(8.65, 4.38 + PH2 + 0.21, 4.45, 0.17, [[("Scope validation → PENDING  ·  AUV hull-slot pod → NEXT BUILD", {"size": 7.6, "italic": True, "color": GRAY})]])

# =========================== TECHNOLOGIES TO BE USED ===========================
TY2 = 5.92
rect(0.30, TY2, 12.80, 0.56, fill=WHITE, line=RULE, lw=0.75)
text(0.40, TY2 + 0.09, 1.30, 0.4, [[("TECHNOLOGIES", {"bold": True, "size": 8.6, "color": NAVY})], [("TO BE USED", {"bold": True, "size": 8.6, "color": NAVY})]])
seg(1.75, TY2 + 0.08, 1.75, TY2 + 0.48, color=RULE, lw=0.75)
techs = [("Embedded", "STM32F446RE · C/C++ · ADC · TIM6 · DMA · DAC"), ("Signal processing", "LFM · Geometric · Barker-13 · FFT · Windowing"),
         ("Analog", "Sallen-Key LPF · TLV9062"), ("Software", "React · Node.js · WebSocket · Python · OpenSCAD")]
for k, (a_, b_) in enumerate(techs):
    cx, cy = 1.90 + (k % 2) * 5.6, TY2 + 0.07 + (k // 2) * 0.23
    text(cx, cy, 5.5, 0.21, [[(a_ + "   ", {"bold": True, "size": 8.4, "color": NAVY}), (b_, {"size": 8.4, "color": CHAR})]])

# =========================== flow + legend ===========================
flow = ["SENSE", "DECIDE", "SYNTHESISE", "STREAM", "CONDITION", "VERIFY"]
widths = [0.78, 0.82, 1.08, 0.86, 1.04, 0.82]; gap = 0.42; x = 0.55
for k, (t, w) in enumerate(zip(flow, widths)):
    rect(x, 6.60, w, 0.24, fill=WHITE, line=LAV if t == "CONDITION" else NAVY, lw=0.75, dashed=(t == "CONDITION"))
    text(x, 6.60, w, 0.24, [[(t, {"bold": True, "size": 8.4, "color": NAVY})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if k < len(flow) - 1: seg(x + w + 0.06, 6.72, x + w + gap - 0.06, 6.72, color=NAVY, lw=0.9, head=True)
    x += w + gap
seg(9.70, 6.72, 10.05, 6.72, color=NAVY, lw=1.0); text(10.12, 6.63, 1.2, 0.18, [[("built / running", {"size": 7.8, "color": GRAY})]])
seg(11.25, 6.72, 11.60, 6.72, color=LAV, lw=1.0, dashed=True); text(11.67, 6.63, 1.45, 0.18, [[("designed / building next", {"size": 7.8, "color": GRAY})]])

prs.save(OUT)
print("wrote", OUT)
