#!/usr/bin/env python3
"""Apply the evidence audit to the team's Canva-exported SIH deck. Reads the original, writes a new file next to it."""
import copy, pathlib, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
SRC = pathlib.Path("/Users/sudharsan/Downloads/SIH2026-IDEA-Presentation-Format.pptx.pptx")
OUT = pathlib.Path("/Users/sudharsan/Downloads/SIH2026-IDEA-AquaSDR-audited.pptx")
HERE = pathlib.Path(__file__).parent
# Fill these before submission; a link left as None is omitted from slide 6 (no visible placeholder).
LINKS = {"Demo video (YouTube)": None, "GitHub repository": "github.com/SIH-Lab/AQUASDR-AUV-Sonar-payload", "Live web console": None}
TEAM_ID, TEAM_NAME = None, None
prs = Presentation(SRC)
NAVY, GRAY, INK, TEAL, LAV = RGBColor(0x18, 0x31, 0x53), RGBColor(0x5B, 0x64, 0x74), RGBColor(0x26, 0x32, 0x38), RGBColor(0x3B, 0x8C, 0x7A), RGBColor(0x7A, 0x6B, 0xA8)

def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6: yield from walk(sh.shapes)
        else: yield sh
def find(slide, name=None, contains=None):
    for sh in walk(slide.shapes):
        if name and sh.name == name: return sh
        if contains and sh.has_text_frame and contains in sh.text_frame.text: return sh
    raise KeyError(name or contains)

def rewrite(shape, items, spacer=False, numbered=False):
    """items: list of paragraphs, each a list of (text, bold). Keeps the box's own paragraph and run formatting."""
    txb = shape.text_frame._txBody
    paras = txb.findall(qn("a:p"))
    tp = next(p for p in paras if p.findall(qn("a:r")))
    runs = [r for p in paras for r in p.findall(qn("a:r"))]
    def rpr(bold):
        for r in runs:
            pr = r.find(qn("a:rPr"))
            if pr is not None and (pr.get("b") == "1") == bold: return copy.deepcopy(pr)
        pr = copy.deepcopy(runs[0].find(qn("a:rPr"))); pr.set("b", "1" if bold else "0")
        lat = pr.find(qn("a:latin"))
        if lat is not None: lat.set("typeface", lat.get("typeface").replace(" Bold", "") + ("" if not bold else " Bold") if bold else lat.get("typeface").replace(" Bold", ""))
        return pr
    RB, RN = rpr(True), rpr(False)
    for p in paras: txb.remove(p)
    seq = []
    for i, it in enumerate(items):
        if spacer and i: seq.append(None)
        seq.append(it)
    num = 0
    for it in seq:
        p = copy.deepcopy(tp)
        if numbered and it:
            num += 1; it = [(f"{num}. ", True)] + list(it)
            ppr = p.find(qn("a:pPr"))
            if ppr is not None:
                for b in list(ppr):
                    if b.tag in (qn("a:buAutoNum"), qn("a:buChar"), qn("a:buFont"), qn("a:buSzPct"), qn("a:buClr"), qn("a:buNone")): ppr.remove(b)
                ppr.append(ppr.makeelement(qn("a:buNone"), {}))
                ppr.set("indent", "0")
        for ch in list(p):
            if ch.tag in (qn("a:r"), qn("a:br"), qn("a:fld")): p.remove(ch)
        end = p.find(qn("a:endParaRPr"))
        for text, bold in (it or []):
            r = p.makeelement(qn("a:r"), {}); r.append(copy.deepcopy(RB if bold else RN))
            tt = r.makeelement(qn("a:t"), {}); tt.text = text; r.append(tt)
            if end is not None: end.addprevious(r)
            else: p.append(r)
        txb.append(p)

def add_text(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, font="DM Sans"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, side, 0)
    for k, para in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); p.alignment = align
        runs, sp = (para[0], para[1]) if isinstance(para, tuple) else (para, 0)
        p.space_after = Pt(sp)
        for text, size, color, bold in runs:
            r = p.add_run(); r.text = text; f = r.font; f.name = font; f.size = Pt(size); f.color.rgb = color; f.bold = bold
    return tb

def swap_image(slide, shape, path):
    blip = next(shape._element.iter(qn("a:blip")))
    _, rid = slide.part.get_or_add_image_part(str(path))
    blip.set(qn("r:embed"), rid)
    bf = blip.getparent()
    src = bf.find(qn("a:srcRect"))
    if src is not None: bf.remove(src)
    for fr in bf.iter(qn("a:fillRect")):
        for k in ("l", "t", "r", "b"): fr.attrib.pop(k, None)
    from PIL import Image
    a = Image.open(path).width / Image.open(path).height
    new_h = int(shape.width / a); shape.top = int(shape.top + (shape.height - new_h) / 2); shape.height = new_h

S = prs.slides
# ---------------- slide 1: title page ----------------
s = S[0]
bl = find(s, "TextBox 9")
for p in bl.text_frame.paragraphs:
    if p.runs and p.runs[0].text.strip().startswith("Problem Statement ID"):
        p.runs[0].text = "Problem Statement ID – 26058"; [r._r.getparent().remove(r._r) for r in p.runs[1:]]
    fill = {"Problem Statement Title": "Problem Statement Title – Development of a Low-Power, Real-Time Adaptive Software-Defined Sonar Transmitter Payload for Autonomous Underwater Vehicles (AUVs)",
            "Theme": "Theme – Robotics and Drones"}
    if TEAM_ID: fill["Team ID"] = f"Team ID – {TEAM_ID}"
    if TEAM_NAME: fill["Team Name"] = f"Team Name – {TEAM_NAME}"
    if p.runs:
        for k, v in fill.items():
            if p.runs[0].text.strip().startswith(k):
                p.runs[0].text = v; [r._r.getparent().remove(r._r) for r in p.runs[1:]]
                if k == "Problem Statement Title": p.runs[0].font.size = Pt(17); p.line_spacing = 1.0
add_text(s, 0.75, 2.62, 8.7, 2.0, [
    ([("AquaSDR", 46, NAVY, True)], 0),
    ([("Adaptive software-defined sonar transmitter for AUVs", 24, NAVY, False)], 6),
    ([("Adapts the next sonar ping to changing water conditions on an STM32, then checks the generated electrical waveform on-board.", 17, GRAY, False)], 0)])

# ---------------- slide 2: idea / solution ----------------
s = S[1]
sub = find(s, contains="Self-Validating Adaptive Sonar")
for p in sub.text_frame.paragraphs:
    for r in p.runs:
        r.text = r.text.replace("Self-Validating Adaptive Sonar Transmitter for AUVs", "Adaptive Software-Defined Sonar Transmitter for AUVs")
rewrite(find(s, "TextBox 53"), [
    [("Sense", True), (": payload sensors read the water; the current adaptive path uses the ", False), ("A0 TDS / salinity proxy", True), (".", False)],
    [("Adapt", True), (": each qualified change re-tunes the ", False), ("next ping", True), (": centre frequency, bandwidth, pulse duration and amplitude.", False)],
    [("Synthesise & stream", True), (": ", False), ("LFM, geometric and Barker-13", True), (" with Hann / Hamming / Blackman windows; TIM6 + DMA → 12-bit DAC (PA4).", False)],
    [("Condition (designed)", True), (": 70.7 kHz low-pass + TLV9062 ×1.5; transducer next; CAD-designed enclosure.", False)]], spacer=True)
rewrite(find(s, "TextBox 46"), [
    [("Fixed ping problem", True), (": one fixed waveform cannot suit both clear and murky water (range vs resolution).", False)],
    [("Water condition changes → next ping reconfigured", True), (": CLEAR 42 kHz, 4 kHz BW, 2 ms, 35 % DAC amplitude → MURKY 38 kHz, 1 kHz BW, 8 ms, 62 %.", False)],
    [("Stable decisions", True), (": EMA (α = 0.25) + hysteresis; a qualified change is applied at the ", False), ("next pulse boundary", True), (".", False)],
    [("No per-sample CPU work", True), (": TIM6 + DMA at 1 MS/s (configured); next pulse built in a spare buffer; ", False), ("0 underruns", True), (" in recorded runs.", False)]], spacer=True)
rewrite(find(s, "TextBox 54"), [
    [("Core idea: measure → qualify → precompute → apply", True), (". The next ping is rebuilt in advance and committed only at the pulse boundary, without disturbing the current one.", False)],
    [("Electrical self-check", True), (": the board samples its own DAC output: ", False), ("42.00 kHz captured vs 42.00 calculated", True), ("; Barker-13 / RECT sidelobe −21.4 vs −21.7 dB.", False)],
    [("Hann window", True), (": measured out-of-band leakage −14.6 → −35 dB (MCU ADC capture).", False)],
    [("Calculated burst-energy ratio", True), (": CLEAR ≈ 12× lower than MURKY.", False)]], spacer=True)

# ---------------- slide 3: technical approach ----------------
s = S[2]
swap_image(s, find(s, "Freeform 18"), HERE / "AquaSDR_slide3_six_panel.png")
swap_image(s, find(s, "Freeform 20"), HERE / "AquaSDR_tech_table_compact.png")
note = find(s, "TextBox 25")
note.left, note.width = Inches(10.99), Inches(8.31)
rewrite(note, [[("Transmit engine validated electrically", True), (" (MCU ADC capture) · analog stage and transducer next", False)]])
for p in note.text_frame.paragraphs: p.alignment = PP_ALIGN.CENTER

# ---------------- slide 4: feasibility ----------------
s = S[3]
rewrite(find(s, "TextBox 30"), [
    [("Built", True), (": STM32F446RE engine streams LFM, geometric and Barker-13 via TIM6 → DMA → DAC (PA4).", False)],
    [("MCU ADC capture", True), (": ", False), ("42.00 kHz", True), (" measured centre and ~4 kHz sweep; matches the calculated reference (13 captures).", False)],
    [("Recorded test runs", True), (": ", False), ("330 / 330", True), (" commands acknowledged; ", False), ("0 DMA underruns", True), (" across the captures.", False)]], numbered=True)
for sh in [x for x in s.shapes if x.name in ("Group 40", "Freeform 43", "TextBox 36")]:   # decorative centre icon and its label
    sh._element.getparent().remove(sh._element)
from pptx.enum.shapes import MSO_SHAPE
bg = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.35), Inches(6.05), Inches(5.3), Inches(3.75))
bg.fill.solid(); bg.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF); bg.line.color.rgb = NAVY; bg.line.width = Pt(1.25); bg.adjustments[0] = 0.08
st = bg._element.find(qn("p:style"))
if st is not None: bg._element.remove(st)
add_text(s, 7.35, 5.72, 5.3, 0.32, [[("▲ PROVEN NOW", 19, NAVY, True)]], align=PP_ALIGN.CENTER)
road = [("TODAY", "38–42 kHz · built and captured", NAVY), ("NEXT", "100–200 kHz electrical validation", TEAL), ("THEN", "200–500 kHz external-DAC path", LAV)]
for i, (k, v, c) in enumerate(road):
    y = 6.30 + i * 1.12
    add_text(s, 7.65, y, 4.8, 1.05, [([(k, 27, c, True)], 2), [(v, 20, INK, False)]])
for nm, txt in (("TextBox 37", "PS gap"), ("TextBox 38", "Next hardware step")):
    sh = find(s, nm); r0 = sh.text_frame.paragraphs[0].runs[0]; r0.text = txt
    for r in sh.text_frame.paragraphs[0].runs[1:]: r._r.getparent().remove(r._r)
rewrite(find(s, "TextBox 32"), [   # left box sits under "Challenge / Risk"
    [("Frequency scale", True), (": current bench prototype 38–42 kHz; PS target 100–500 kHz.", False)],
    [("Analog chain not built", True), (": low-pass filter, TLV9062 stage and transducer are designed only.", False)],
    [("External proof pending", True), (": oscilloscope capture, power measurement and an underwater tank test.", False)]], numbered=True)
rewrite(find(s, "TextBox 31"), [   # right box sits under "Strategies"
    [("PS-band path", True), (": 100–200 kHz is the next electrical validation target; 200–500 kHz is a planned external-DAC architecture (≥ 2.5 MS/s).", False)],
    [("Build the designed stage", True), (": 70.7 kHz Sallen-Key + TLV9062 ×1.5; parts list and SPICE netlist ready.", False)],
    [("Validate in order", True), (": scope capture → power meter → tank test with a transducer.", False)]], numbered=True)

# ---------------- slide 5: impact ----------------
s = S[4]
for nm in ("TextBox 6", "TextBox 7", "TextBox 8", "TextBox 9", "TextBox 10", "TextBox 11", "TextBox 12"):   # hidden behind the cards
    sh = find(s, nm); sh._element.getparent().remove(sh._element)
rewrite(find(s, "TextBox 31"), [
    [("Research & field teams", True), (": flexible transmit profiles for changing water conditions.", False)],
    [("Target use", True), (": research and inspection missions.", False)]], numbered=True)
rewrite(find(s, "TextBox 32"), [
    [("One configurable digital engine", True), (": new profiles through firmware, not a new waveform-generation hardware path.", False)],
    [("Low-cost hardware", True), (": off-the-shelf STM32, op-amp and sensors.", False)]], numbered=True)
rewrite(find(s, "TextBox 33"), [
    [("Energy-aware transmission", True), (": lower-amplitude profile where the selected condition permits (35 % vs 62 % DAC amplitude).", False)],
    [("~12× lower burst energy", True), (" for CLEAR vs MURKY (calculated); acoustic impact to be measured.", False)]], numbered=True)
sh = find(s, "TextBox 37"); sh.text_frame.paragraphs[0].runs[0].text = "Target users"

# ---------------- slide 6: research and references ----------------
s = S[5]
add_text(s, 0.9, 2.1, 8.6, 0.5, [[("PROJECT LINKS", 26, NAVY, True)]])
shown = [(k, v) for k, v in LINKS.items() if v]
for i, (a, b) in enumerate(shown):
    y = 2.85 + i * 1.05
    add_text(s, 0.9, y, 8.8, 0.95, [([(a, 20, NAVY, True)], 2), [(b, 19, TEAL, False)]])
LY = 2.85 + len(shown) * 1.05 + 0.35
add_text(s, 0.9, LY, 8.6, 0.5, [[("EVIDENCE LEGEND", 26, NAVY, True)]])
legend = [("CALCULATED", "from the C reference engine (deterministic)"), ("MCU ADC CAPTURE", "the board samples its own raw DAC output"),
          ("DESIGNED", "modelled and simulated, not yet built"), ("PENDING", "external oscilloscope, power and acoustic tests")]
for i, (a, b) in enumerate(legend):
    add_text(s, 0.9, LY + 0.7 + i * 0.82, 8.8, 0.8, [([(a, 17, [NAVY, TEAL, LAV, GRAY][i], True)], 1), [(b, 16, INK, False)]])
add_text(s, 10.3, 2.1, 9.0, 0.5, [[("REFERENCES", 26, NAVY, True)]])
refs = ["NIOT / Ministry of Earth Sciences, SIH 2026 Problem Statement 26058.",
        "F. J. Harris, “On the use of windows for harmonic analysis with the DFT,” Proc. IEEE, 66(1), 1978.",
        "R. H. Barker, “Group synchronizing of binary digital systems,” in Communication Theory, 1953.",
        "M. I. Skolnik, Introduction to Radar Systems, 3rd ed., McGraw-Hill, 2001 (LFM pulse compression).",
        "X. Lurton, An Introduction to Underwater Acoustics, 2nd ed., Springer, 2010.",
        "STMicroelectronics, RM0390 STM32F446xx reference manual; AN3126 DAC waveform generation.",
        "Texas Instruments, TLV9062 datasheet; SLOA024 Analysis of the Sallen-Key architecture.",
        "Sensor datasheets: DS18B20, BMP390, INA219, ISM330DHCX, MMC5983MA.",
        "Ministry of Earth Sciences, Deep Ocean Mission."]
add_text(s, 10.3, 2.85, 9.0, 7.3, [([(f"{i}.  ", 16.5, NAVY, True), (r, 16.5, INK, False)], 9) for i, r in enumerate(refs, 1)])

prs.save(OUT); print("wrote", OUT)
