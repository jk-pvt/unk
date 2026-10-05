#!/usr/bin/env python3
"""SIH 2026 deck, slide 4 (FEASIBILITY AND VIABILITY) in the slide-3 v5 style.
Left: ANALYSIS OF FEASIBILITY: a readiness matrix (subsystem x maturity stage) plus six proof numbers, all from the
repository's recorded evidence (data/bench MCU ADC captures, docs/evidence). Right: CHALLENGES & RISKS -> STRATEGIES.
Bottom: VIABILITY line. Honest maturity lives here (slide 3 shows the architecture only).
    python3 docs/deck/build_feasibility_slide.py                 -> full slide on the SIH template
    FIGURE_ONLY=1 python3 docs/deck/build_feasibility_slide.py   -> figure only (for pasting into Canva)
"""
import os, pathlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn

HERE = pathlib.Path(__file__).parent
FIGURE_ONLY = os.environ.get("FIGURE_ONLY") == "1"
OUT = HERE / ("AquaSDR_slide4_figure.pptx" if FIGURE_ONLY else "AquaSDR_slide4_feasibility.pptx")
rgb = RGBColor.from_string
NAVY, TEAL, BG, LAV, CHAR, GRAY, WHITE, RULE, FAINT = rgb("183153"), rgb("3B8C7A"), rgb("F3F6F9"), rgb("8A7BAF"), rgb("263238"), rgb("667085"), rgb("FFFFFF"), rgb("D0D5DD"), rgb("C3C9D3")
FONT = "Arial"

prs = Presentation(HERE / "template" / "SIH2026-IDEA-template.pptx")
ids = prs.slides._sldIdLst
for i, sid in reversed(list(enumerate(list(ids)))):
    if i != 3: prs.part.drop_rel(sid.rId); ids.remove(sid)
s = prs.slides[0]
for sh in list(s.shapes):
    if sh.has_text_frame and "Analysis of the feasibility" in sh.text_frame.text:
        sh._element.getparent().remove(sh._element)
    elif sh.has_text_frame and sh.text_frame.text.strip() == "FEASIBILITY AND VIABILITY":
        p = sh.text_frame.paragraphs[0]
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
        p.runs[0].text = "AquaSDR · FEASIBILITY AND VIABILITY"; p.runs[0].font.size = Pt(28)
if FIGURE_ONLY:
    for sh in list(s.shapes): sh._element.getparent().remove(sh._element)

I = Inches
def flat(shape):
    el = shape._element
    st = el.find(qn("p:style"))
    if st is not None: el.remove(st)
    spPr = el.find(qn("p:spPr"))
    if spPr is not None and spPr.find(qn("a:effectLst")) is None: spPr.append(spPr.makeelement(qn("a:effectLst"), {}))
    return shape

def text(x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, size=8.5, color=CHAR, bold=False, italic=False):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, side, 0)
    if isinstance(paras, str): paras = [paras]
    for k, para in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); p.alignment = align
        for t, st in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run(); r.text = t; f = r.font; f.name = FONT
            f.size = Pt(st.get("size", size)); f.bold = st.get("bold", bold); f.italic = st.get("italic", italic); f.color.rgb = st.get("color", color)
    return tb

def rect(x, y, w, h, fill=None, line=NAVY, lw=0.75, dashed=False):
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
    if fill is None: b.fill.background()
    else: b.fill.solid(); b.fill.fore_color.rgb = fill
    if line is None: b.line.fill.background()
    else:
        b.line.color.rgb = line; b.line.width = Pt(lw)
        if dashed: b.line.dash_style = MSO_LINE.DASH
    return flat(b)

def seg(x1, y1, x2, y2, color=NAVY, lw=1.0, head=False, dashed=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = color; c.line.width = Pt(lw); flat(c)
    if dashed: c.line.dash_style = MSO_LINE.DASH
    if head:
        ln = c.line._get_or_add_ln(); ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "sm", "len": "sm"}))
    return c

def circle(cx, cy, r, fill=None, line=NAVY, lw=1.0):
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, I(cx - r), I(cy - r), I(2 * r), I(2 * r))
    if fill is None: d.fill.background()
    else: d.fill.solid(); d.fill.fore_color.rgb = fill
    if line is None: d.line.fill.background()
    else: d.line.color.rgb = line; d.line.width = Pt(lw)
    return flat(d)

def marker(cx, cy, state):
    """done = filled navy · part = navy ring with a centre dot · todo = faint ring · na = n/a"""
    if state == "done": circle(cx, cy, 0.075, fill=NAVY, line=NAVY)
    elif state == "part": circle(cx, cy, 0.075, fill=WHITE, line=NAVY, lw=1.2); circle(cx, cy, 0.032, fill=NAVY, line=None)
    elif state == "todo": circle(cx, cy, 0.075, fill=WHITE, line=FAINT, lw=1.0)
    else: text(cx - 0.2, cy - 0.08, 0.4, 0.16, [[("n/a", {"size": 6.8, "color": FAINT})]], align=PP_ALIGN.CENTER)

def header(x, y, w, label, subtitle):
    text(x, y, w, 0.18, [[(label, {"bold": True, "size": 9.6, "color": NAVY})]])
    text(x, y + 0.18, w, 0.17, [[(subtitle, {"size": 8.6, "color": GRAY})]])
    seg(x, y + 0.38, x + w, y + 0.38, color=RULE, lw=0.75)

if not FIGURE_ONLY:
    text(2.0, 0.93, 9.3, 0.24, [[("Proven on our hardware  ·  risks identified  ·  a mitigation for each", {"size": 12, "italic": True, "color": GRAY})]], align=PP_ALIGN.CENTER)

# =========================== ANALYSIS OF FEASIBILITY ===========================
header(0.30, 1.30, 6.30, "ANALYSIS OF FEASIBILITY", "Readiness of every subsystem, from design to instrument validation")
MX, MY, LWID, CWID, RH = 0.30, 1.80, 2.05, 0.85, 0.285
stages = ["DESIGNED", "SIMULATED\n/ TESTED", "BUILT", "ON-BOARD\nVERIFIED", "INSTRUMENT\nVALIDATED"]
rows = [("Adaptive firmware engine", "C/C++ on STM32F446RE", ["done", "done", "done", "done", "todo"]),
        ("Timer + DMA + DAC streaming", "TIM6 → DMA1 → DAC1 · 1 MS/s", ["done", "done", "done", "done", "todo"]),
        ("LFM · geometric · Barker-13", "with Hann / Hamming / Blackman", ["done", "done", "done", "done", "todo"]),
        ("Sensor input + adaptation", "CLEAR / TRANSITION / MURKY", ["done", "done", "done", "part", "todo"]),
        ("Analog front end", "Sallen-Key LPF + TLV9062", ["done", "done", "todo", "todo", "todo"]),
        ("Payload enclosure", "acrylic hull · printed AUV pod", ["done", "done", "part", "na", "na"]),
        ("Low-power operation", "DMA-driven, CPU free per sample", ["done", "done", "na", "todo", "todo"])]
HY = MY + 0.36
rect(MX, MY, LWID + 5 * CWID, 0.36, fill=BG, line=None)
for j, st in enumerate(stages):
    text(MX + LWID + j * CWID, MY + 0.03, CWID, 0.32, [[(line, {"bold": True, "size": 6.6, "color": NAVY})] for line in st.split("\n")], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
seg(MX, HY, MX + LWID + 5 * CWID, HY, color=NAVY, lw=0.75)
for i, (name, sub, states) in enumerate(rows):
    y = HY + i * RH
    text(MX + 0.04, y + 0.015, LWID - 0.06, 0.15, [[(name, {"bold": True, "size": 7.8, "color": NAVY})]])
    text(MX + 0.04, y + 0.145, LWID - 0.06, 0.14, [[(sub, {"size": 6.6, "color": GRAY})]])
    for j, stt in enumerate(states):
        cx = MX + LWID + j * CWID + CWID / 2
        if j < 4 and states[j] in ("done", "part") and states[j + 1] in ("done", "part"):
            seg(cx + 0.075, y + RH / 2, cx + CWID - 0.075, y + RH / 2, color=NAVY, lw=1.25)   # progress bar between reached stages
        marker(cx, y + RH / 2, stt)
    seg(MX, y + RH, MX + LWID + 5 * CWID, y + RH, color=RULE, lw=0.5)
LY = HY + len(rows) * RH + 0.05
lx = MX
for st, lab in [("done", "reached"), ("part", "partly (A0 shown in CLEAR; printed pod next)"), ("todo", "next")]:
    marker(lx + 0.08, LY + 0.08, st); text(lx + 0.2, LY + 0.005, 3.2, 0.16, [[(lab, {"size": 6.8, "color": GRAY})]]); lx += 0.85 if st != "part" else 2.75

# proof numbers (real recorded evidence)
PY0 = LY + 0.27
text(MX, PY0, 6.3, 0.17, [[("PROVEN ON THE BOARD ", {"bold": True, "size": 8.2, "color": NAVY}), ("· MCU ADC capture and firmware telemetry, not oscilloscope data", {"size": 7.2, "color": GRAY})]])
tiles = [("42.00 kHz", "captured centre = calculated", "LFM · Hann"), ("−14.6 → −35 dB", "spectral leakage, rect → Hann", "window effect captured"),
         ("−21.4 dB", "Barker-13 sidelobe captured", "−21.7 dB calculated"), ("330 / 330", "commands acknowledged", "0 lost"),
         ("0", "DMA underruns", "across every capture"), ("< 100 ms", "decision → new waveform", "86–97 ms measured")]
TW, TH = 2.05, 0.58
for k, (big, a, b) in enumerate(tiles):
    tx, ty = MX + (k % 3) * (TW + 0.075), PY0 + 0.22 + (k // 3) * (TH + 0.06)
    rect(tx, ty, TW, TH, fill=WHITE, line=NAVY, lw=0.75)
    rect(tx, ty, 0.05, TH, fill=TEAL, line=None)
    text(tx + 0.14, ty + 0.03, TW - 0.2, 0.26, [[(big, {"bold": True, "size": 13, "color": NAVY})]])
    text(tx + 0.14, ty + 0.285, TW - 0.2, 0.15, [[(a, {"size": 7.2, "color": CHAR})]])
    text(tx + 0.14, ty + 0.42, TW - 0.2, 0.15, [[(b, {"size": 6.6, "color": GRAY})]])

# =========================== CHALLENGES → STRATEGIES ===========================
RX, RW = 6.85, 6.25
header(RX, 1.30, 4.75, "CHALLENGES & RISKS  →  STRATEGIES", "Each risk paired with how we overcome it")
CW1, CW2 = 2.30, 2.86
risks = [("Frequency scale", "PS cites 100–500 kHz; bench transducer is 40 kHz",
          "Frequency-agnostic policy table; external high-speed DAC on the same TIM + DMA path", "SCALE PATH"),
         ("No oscilloscope on hand", "external validation not yet possible",
          "On-board ADC + DMA capture with FFT and matched filter now; lab scope at judging, CSV import ready", "IN PLACE"),
         ("Analog stage not yet built", "filter + op-amp still to assemble",
          "Two independent models agree within 0.02 dB in band; stock-value parts list ready", "NEXT BUILD"),
         ("Noisy or drifting sensor input", "a bad reading could change the ping",
          "Smoothing, hysteresis, stale and out-of-range rejection; TDS labelled as a proxy", "IN PLACE"),
         ("CPU load and battery budget", "per-sample trig math is expensive",
          "Next pulse computed between pings, DMA streams it; inline-meter power test planned", "IN PLACE"),
         ("Loss of control link", "ping must never run unattended",
          "Output off at boot, 2 s control lease auto-stop, DAC parked mid-scale on any fault", "IN PLACE"),
         ("Water ingress", "payload must survive the hull slot",
          "O-ring end caps; parametric pod compiled and fit-checked in OpenSCAD", "DESIGNED")]
CY0, CRH = 1.80, 0.555
rect(RX, CY0, RW, 0.26, fill=BG, line=None)
text(RX + 0.08, CY0, CW1, 0.26, [[("CHALLENGE", {"bold": True, "size": 7.2, "color": NAVY})]], anchor=MSO_ANCHOR.MIDDLE)
text(RX + CW1 + 0.35, CY0, CW2, 0.26, [[("STRATEGY", {"bold": True, "size": 7.2, "color": NAVY})]], anchor=MSO_ANCHOR.MIDDLE)
text(RX + RW - 0.72, CY0, 0.7, 0.26, [[("STATUS", {"bold": True, "size": 7.2, "color": NAVY})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
seg(RX, CY0 + 0.26, RX + RW, CY0 + 0.26, color=NAVY, lw=0.75)
for i, (c, cd, strat, status) in enumerate(risks):
    y = CY0 + 0.26 + i * CRH
    text(RX + 0.08, y + 0.07, CW1 - 0.1, 0.17, [[(c, {"bold": True, "size": 8.2, "color": NAVY})]])
    text(RX + 0.08, y + 0.25, CW1 - 0.1, 0.28, [[(cd, {"size": 7.0, "color": GRAY})]])
    seg(RX + CW1 + 0.05, y + CRH / 2, RX + CW1 + 0.27, y + CRH / 2, color=TEAL, lw=1.25, head=True)
    text(RX + CW1 + 0.35, y + 0.08, CW2, 0.42, [[(strat, {"size": 7.8, "color": CHAR})]], anchor=MSO_ANCHOR.MIDDLE)
    done = status == "IN PLACE"
    rect(RX + RW - 0.70, y + 0.17, 0.66, 0.21, fill=WHITE, line=NAVY if done else LAV, lw=0.75, dashed=not done)
    text(RX + RW - 0.70, y + 0.17, 0.66, 0.21, [[(status, {"bold": True, "size": 6.0, "color": NAVY if done else LAV})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    seg(RX, y + CRH, RX + RW, y + CRH, color=RULE, lw=0.5)

# =========================== VIABILITY ===========================
VY = 6.05
rect(0.30, VY, 12.80, 0.48, fill=WHITE, line=RULE, lw=0.75)
text(0.40, VY, 1.2, 0.48, [[("VIABILITY", {"bold": True, "size": 9, "color": NAVY})]], anchor=MSO_ANCHOR.MIDDLE)
seg(1.45, VY + 0.08, 1.45, VY + 0.40, color=RULE, lw=0.75)
pts = [("Low-cost COTS", "STM32 Nucleo · TLV9062 · standard passives"), ("Software-defined", "new waveform or policy = firmware update"),
       ("Scalable", "same architecture → 100–500 kHz with an external DAC"), ("Maintainable", "161 automated tests guard every change")]
for k, (a, b) in enumerate(pts):
    cx = 1.58 + k * 2.88
    text(cx, VY + 0.07, 2.8, 0.36, [[(a, {"bold": True, "size": 8.0, "color": NAVY})], [(b, {"size": 7.2, "color": CHAR})]])

prs.save(OUT)
print("wrote", OUT)
