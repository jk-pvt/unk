#!/usr/bin/env python3
"""SIH 2026 deck, slide 3 (TECHNICAL APPROACH): AquaSDR system architecture built as native, editable PowerPoint
objects on the official SIH template, with the team's own hardware photographs (docs/deck/img, from the user's photos).

Status language (two styles only): solid border = built / running; dashed border = designed / next build.
Every number on the slide comes from the repository's recorded evidence (docs/evidence, data/bench).
    python3 docs/deck/build_technical_slide.py      -> docs/deck/AquaSDR_slide3_technical_approach.pptx
"""
import copy, pathlib
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn

HERE = pathlib.Path(__file__).parent
IMG = HERE / "img"
OUT = HERE / "AquaSDR_slide3_technical_approach.pptx"

NAVY, SUB = RGBColor(0x14, 0x21, 0x3D), RGBColor(0x4A, 0x55, 0x68)
TEAL, BLUE, LAV, RED, BEIGE = RGBColor(0x2E, 0x8B, 0x6E), RGBColor(0x2A, 0x4D, 0x8F), RGBColor(0x7B, 0x5E, 0xA7), RGBColor(0xB0, 0x3A, 0x2E), RGBColor(0x9A, 0x6A, 0x2F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

prs = Presentation(HERE / "template" / "SIH2026-IDEA-template.pptx")
# keep only the TECHNICAL APPROACH slide (index 2); the others stay in the user's own deck
ids = prs.slides._sldIdLst
for i, sid in reversed(list(enumerate(list(ids)))):
    if i != 2:
        prs.part.drop_rel(sid.rId); ids.remove(sid)
s = prs.slides[0]
for sh in list(s.shapes):                 # the template's pointer text box is re-used as section labels on the figure
    if sh.has_text_frame and "Technologies to be used" in sh.text_frame.text:
        sh._element.getparent().remove(sh._element)

I = Inches
def text(x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, size=9, color=NAVY, bold=False, italic=False):
    """paras: str | list of paragraphs, each a str or a list of (text, {size,bold,color,italic}) runs."""
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, side, 0)
    if isinstance(paras, str): paras = [paras]
    for k, para in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); p.alignment = align
        runs = [(para, {})] if isinstance(para, str) else para
        for t, st in runs:
            r = p.add_run(); r.text = t; f = r.font; f.name = FONT
            f.size = Pt(st.get("size", size)); f.bold = st.get("bold", bold); f.italic = st.get("italic", italic)
            f.color.rgb = st.get("color", color)
    return tb

def box(x, y, w, h, fill=None, line=BLUE, lw=1.25, dashed=False, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    b = s.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: b.adjustments[0] = radius
    if fill is None: b.fill.background()
    else: b.fill.solid(); b.fill.fore_color.rgb = fill
    if line is None: b.line.fill.background()
    else:
        b.line.color.rgb = line; b.line.width = Pt(lw)
        if dashed: b.line.dash_style = MSO_LINE.DASH
    b.shadow.inherit = False
    b.text_frame.text = ""
    return b

def seg(x1, y1, x2, y2, color=NAVY, lw=1.5, head=False, dashed=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = color; c.line.width = Pt(lw)
    if dashed: c.line.dash_style = MSO_LINE.DASH
    if head:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c

def path(points, color=NAVY, lw=1.5):           # polyline, arrow on the last segment
    for k in range(len(points) - 1):
        seg(*points[k], *points[k + 1], color=color, lw=lw, head=(k == len(points) - 2))

def dot(x, y, r=0.045, fill=WHITE, line=TEAL, lw=1.25):
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, I(x - r), I(y - r), I(2 * r), I(2 * r))
    d.fill.solid(); d.fill.fore_color.rgb = fill; d.line.color.rgb = line; d.line.width = Pt(lw); d.shadow.inherit = False
    return d

def badge(x, y, num, color):
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, I(x), I(y), I(0.27), I(0.27))
    d.fill.solid(); d.fill.fore_color.rgb = color; d.line.fill.background(); d.shadow.inherit = False
    tf = d.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = num
    r.font.name = FONT; r.font.size = Pt(8); r.font.bold = True; r.font.color.rgb = WHITE

def picture(path_, x, y, w, h, line=None, dashed=False):
    p = s.shapes.add_picture(str(path_), I(x), I(y), I(w), I(h))
    if line is not None:
        p.line.color.rgb = line; p.line.width = Pt(1.25)
        if dashed: p.line.dash_style = MSO_LINE.DASH
    return p

# =============== ZONE A · SENSE (left) ===============
text(0.30, 1.30, 3.3, 0.25, [[("SENSE · REAL ENVIRONMENT", {"bold": True, "color": TEAL, "size": 10.5})]])
PX, PY, PW = 1.45, 1.60, 2.05
PH = PW / 0.909
picture(IMG / "sensors.jpg", PX, PY, PW, PH, line=TEAL)
P = lambda u, v: (PX + u / 1400 * PW, PY + v / 1541 * PH)
callouts = [("TDS / salinity proxy", (514, 424), 1.92), ("HC-SR04 ultrasonic", (719, 617), 2.32),
            ("Turbidity sensor", (540, 822), 2.72), ("DS18B20 temperature", (1079, 796), 3.12), ("Pressure / depth", (1028, 1130), 3.52)]
for label, (u, v), ly in callouts:
    px, py = P(u, v)
    text(0.18, ly - 0.09, 1.22, 0.2, [[(label, {"bold": True, "size": 7.6})]], align=PP_ALIGN.RIGHT)
    seg(1.43, ly, px, py, color=TEAL, lw=1.0); dot(px, py)
text(PX, PY + PH + 0.05, PW, 0.2, [[("ADC  ·  1-Wire  ·  I²C", {"bold": True, "color": TEAL, "size": 9})]], align=PP_ALIGN.CENTER)
text(0.30, 4.12, 3.25, 0.32, [[("also on board: IMU + magnetometer · BMP390 cabin pressure · INA219 ×2 (mounted, not yet commissioned)", {"size": 7.5, "color": SUB})]])
# one policy, two inputs
text(0.30, 4.50, 3.25, 0.2, [[("ONE ADAPTATION POLICY · TWO INPUTS", {"bold": True, "size": 8, "color": NAVY})]])
for k, (lbl, sub) in enumerate([("A0 sensor", "TDS_PROXY · ADC"), ("Web command", "console → UART")]):
    bx = 0.30 + k * 1.50
    box(bx, 4.74, 1.40, 0.44, fill=RGBColor(0xEE, 0xF7, 0xF2), line=TEAL, lw=1.0)
    text(bx, 4.77, 1.40, 0.4, [[(lbl, {"bold": True, "size": 8.5})], [(sub, {"size": 7.5, "color": SUB})]], align=PP_ALIGN.CENTER)
dot(3.40, 4.96, r=0.05, fill=TEAL, line=TEAL)
path([(3.20, 4.96), (3.66, 4.96), (3.66, 2.52), (3.86, 2.52)], color=TEAL, lw=1.5)   # both -> 02 DECIDE
seg(1.70, 5.30, 3.40, 5.30, color=TEAL, lw=1.0); seg(1.00, 5.18, 1.00, 5.30, color=TEAL, lw=1.0); seg(3.40, 5.30, 3.40, 4.96, color=TEAL, lw=1.0)
seg(1.00, 5.30, 1.70, 5.30, color=TEAL, lw=1.0); seg(2.40, 5.18, 2.40, 5.30, color=TEAL, lw=1.0)
# technologies (template pointer 1)
box(0.30, 5.52, 3.25, 0.88, fill=RGBColor(0xF7, 0xF8, 0xFA), line=RGBColor(0xC3, 0xCA, 0xD6), lw=0.75)
text(0.40, 5.56, 3.1, 0.95, [[("TECHNOLOGIES TO BE USED", {"bold": True, "size": 8, "color": NAVY})],
                             [("STM32F446RE (Cortex-M4F) · embedded C / C++", {"size": 7.5, "color": SUB})],
                             [("TIM6 + DMA + 12-bit DAC · ADC / 1-Wire / I²C", {"size": 7.5, "color": SUB})],
                             [("Sallen-Key LPF + TLV9062 op-amp (analog)", {"size": 7.5, "color": SUB})],
                             [("Node.js + React console · Python / NumPy · OpenSCAD", {"size": 7.5, "color": SUB})]])
seg(PX + PW, 2.06, 3.86, 2.06, color=TEAL, lw=1.5, head=True)   # sensors -> 01 ACQUIRE

# =============== ZONE B · STM32 (centre) ===============
text(3.75, 1.30, 4.7, 0.25, [[("STM32F446RE · EMBEDDED CONTROL", {"bold": True, "color": BLUE, "size": 10.5}),
                              ("    methodology & process", {"size": 8, "color": SUB, "italic": True})]])
box(3.75, 1.58, 4.70, 4.72, fill=RGBColor(0xF4, 0xF7, 0xFC), line=BLUE, lw=1.25, radius=0.03)
key = box(3.85, 1.66, 4.50, 0.26, fill=NAVY, line=None, radius=0.25)
text(3.85, 1.66, 4.50, 0.26, [[("THE NEXT PING IS COMPUTED ON THE MCU", {"bold": True, "size": 9.5, "color": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
steps = [("01", "ACQUIRE", "ADC · 1-Wire · I²C · stale / out-of-range readings rejected", BLUE),
         ("02", "DECIDE", "smoothing + hysteresis → CLEAR / TRANSITION / MURKY", BLUE),
         ("03", "SYNTHESISE", "LFM chirp · geometric sweep · Barker-13  ×  Hann / Hamming / Blackman", BLUE),
         ("04", "STREAM", "TIM6 → DMA → 12-bit DAC1 (PA4) · 1 MS/s · CPU sleeps between pings", BLUE),
         ("06", "SELF-VALIDATE", "ADC3 + DMA capture of the DAC pin (PA4 → A5) → FFT / spectral analysis", RED)]
Y0, DY, CX = 1.98, 0.47, 4.03
seg(CX + 0.135, Y0 + 0.13, CX + 0.135, Y0 + 3 * DY + 0.13, color=BLUE, lw=1.5)
for k, (n, t, d, col) in enumerate(steps):
    y = Y0 + k * DY
    badge(CX, y, n, col)
    text(4.40, y - 0.02, 3.95, 0.2, [[(t, {"bold": True, "size": 9.5, "color": NAVY})]])
    text(4.40, y + 0.17, 3.95, 0.22, [[(d, {"size": 7.8, "color": SUB})]])
# self-validation loop: DAC pin back into ADC3
path([(3.92, Y0 + 3 * DY + 0.13), (3.92, Y0 + 4 * DY + 0.13), (4.02, Y0 + 4 * DY + 0.13)], color=RED, lw=1.25)
# evidence strip
box(3.85, 4.38, 4.50, 0.40, fill=WHITE, line=RGBColor(0xC3, 0xCA, 0xD6), lw=0.75, radius=0.15)
text(3.95, 4.40, 4.35, 0.36, [[("MCU ADC CAPTURE  ", {"bold": True, "size": 7.4, "color": RED}),
                               ("42.00 kHz centre · 4 kHz sweep · 2 ms pulse · HANN", {"size": 7.4, "color": NAVY})],
                              [("LINK + STREAM  ", {"bold": True, "size": 7.4, "color": BLUE}),
                               ("330 / 330 commands acknowledged · 0 DMA underruns", {"size": 7.4, "color": NAVY})]],
     anchor=MSO_ANCHOR.MIDDLE)
# real electronics bay photo with callouts
BX, BY, BW = 4.20, 4.84, 3.80
BH = BW / 2.7027
picture(IMG / "bay.jpg", BX, BY, BW, BH, line=BLUE)
for label, (fx, fy), (lx, ly) in [("STM32 NUCLEO-F446RE", (0.278, 0.114), (0.02, 0.62)), ("BMP390", (0.522, 0.895), (0.43, 0.62)),
                                   ("TDS meter (A0)", (0.683, 0.82), (0.74, 0.15))]:
    tx, ty = BX + fx * BW, BY + fy * BH
    w = 0.062 * len(label) + 0.16
    lx_, ly_ = BX + lx * BW, BY + ly * BH
    pill = box(lx_, ly_, w, 0.19, fill=WHITE, line=BLUE, lw=0.75, radius=0.5)
    text(lx_, ly_ + 0.01, w, 0.17, [[(label, {"bold": True, "size": 7, "color": NAVY})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    seg(lx_ + w / 2, ly_ + (0 if ly_ > ty else 0.19), tx, ty, color=WHITE, lw=1.0); dot(tx, ty, fill=WHITE, line=BLUE)

# =============== ZONE C · NEXT PING → CONDITION → PAYLOAD (right) ===============
text(8.65, 1.30, 2.95, 0.25, [[("ADAPTED NEXT PING", {"bold": True, "color": LAV, "size": 10.5})]])
cols = [("CLEAR", RGBColor(0x3A, 0x7B, 0xD5)), ("TRANSITION", LAV), ("MURKY", BEIGE)]
LX, CW, TY = 8.72, 0.69, 1.62
for j, (c, colr) in enumerate(cols):
    cx = LX + 0.86 + j * CW
    box(cx + 0.02, TY, CW - 0.04, 0.22, fill=colr, line=None, radius=0.3)
    text(cx, TY + 0.01, CW, 0.2, [[(c, {"bold": True, "size": 7, "color": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rows = [("Centre", ["42 kHz", "40 kHz", "38 kHz"]), ("Bandwidth", ["4 kHz", "2 kHz", "1 kHz"]),
        ("Pulse", ["2 ms", "4 ms", "8 ms"]), ("Amplitude", ["35 %", "50 %", "62 %"])]
for i, (n, vals) in enumerate(rows):
    y = TY + 0.28 + i * 0.22
    text(LX, y, 0.86, 0.2, [[(n, {"bold": True, "size": 8.5})]])
    for j, v in enumerate(vals): text(LX + 0.86 + j * CW, y, CW, 0.2, [[(v, {"size": 8.5, "color": SUB})]], align=PP_ALIGN.CENTER)
    if i < 3: seg(LX, y + 0.205, LX + 0.86 + 3 * CW, y + 0.205, color=RGBColor(0xE2, 0xE6, 0xEE), lw=0.5)
text(LX, 2.78, 4.45, 0.36, [[("CLEAR → shorter, sharper ping    MURKY → lower band · longer · more output", {"size": 7.8, "italic": True, "color": SUB})]])
path([(8.45, Y0 + DY + 0.13), (8.55, Y0 + DY + 0.13), (8.55, 2.18), (8.63, 2.18)], color=LAV, lw=1.5)     # 02 -> next ping

# 05 CONDITION (designed): dashed
box(8.65, 3.08, 4.45, 1.30, fill=WHITE, line=LAV, lw=1.25, dashed=True, radius=0.05)
badge(8.74, 3.15, "05", LAV)
text(9.07, 3.16, 2.5, 0.25, [[("CONDITION · analog front end", {"bold": True, "size": 9.5, "color": NAVY})]])
picture(IMG / "afe_chain.png", 8.78, 3.50, 2.62, 2.62 / 3.645)
text(11.50, 3.50, 1.55, 0.6, [[("2nd-order LPF", {"bold": True, "size": 8.5})], [("70.7 kHz cut-off", {"size": 8, "color": SUB})], [("TLV9062 ×1.5", {"bold": True, "size": 8.5})]])
box(11.50, 4.08, 1.48, 0.21, fill=WHITE, line=LAV, lw=0.75, dashed=True, radius=0.5)
text(11.50, 4.09, 1.48, 0.19, [[("DESIGNED · NEXT BUILD", {"bold": True, "size": 6.8, "color": LAV})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
seg(8.45, Y0 + 3 * DY + 0.13, 8.65, Y0 + 3 * DY + 0.13, color=LAV, lw=1.5, head=True)            # 04 DAC -> 05

# physical payload (built)
seg(10.875, 4.38, 10.875, 4.60, color=NAVY, lw=1.5, head=True)
picture(IMG / "payload.jpg", 8.65, 4.62, 4.45, 4.45 / 3.735, line=NAVY)
text(8.65, 5.84, 4.45, 0.2, [[("AQUASDR PAYLOAD  ", {"bold": True, "size": 9, "color": NAVY}), ("physical prototype · acrylic payload hull", {"size": 8, "color": SUB})]])
text(8.65, 6.06, 4.45, 0.2, [[("External scope validation → PENDING   ·   AUV hull-slot pod → NEXT BUILD", {"size": 7.5, "color": SUB, "italic": True})]])

# =============== legend + storyline ===============
box(10.70, 6.33, 0.30, 0.15, fill=WHITE, line=NAVY, lw=1.0, radius=0.2); text(11.05, 6.31, 0.9, 0.2, [[("built / running", {"size": 7, "color": SUB})]])
box(11.95, 6.33, 0.30, 0.15, fill=WHITE, line=NAVY, lw=1.0, dashed=True, radius=0.2); text(12.30, 6.31, 1.0, 0.2, [[("designed / next build", {"size": 7, "color": SUB})]])
text(0.30, 6.60, 12.8, 0.28, [[("SENSE THE WATER  →  COMPUTE THE NEXT PING  →  STREAM THE WAVEFORM  →  VERIFY THE OUTPUT", {"bold": True, "size": 11.5, "color": NAVY})]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print("wrote", OUT)
