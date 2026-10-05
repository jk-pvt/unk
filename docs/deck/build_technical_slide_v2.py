#!/usr/bin/env python3
"""SIH 2026 deck, slide 3 (TECHNICAL APPROACH), redesign v2: one continuous engineering story
water -> sensors -> STM32 -> policy -> next ping -> synthesis -> TIM6+DMA+DAC -> analog stage -> payload -> validation,
built as native, editable PowerPoint objects on the official SIH template, using only the team's own photographs.

Conventions: solid border = built / running; dashed border = designed / next build. Numbers come from the repository's
recorded evidence (data/bench, docs/evidence). Nothing on the slide is an oscilloscope, analog-stage or acoustic result.
    python3 docs/deck/build_technical_slide_v2.py   -> docs/deck/AquaSDR_slide3_technical_approach_v2.pptx
"""
import pathlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn

HERE = pathlib.Path(__file__).parent
IMG = HERE / "img"
OUT = HERE / "AquaSDR_slide3_technical_approach_v2.pptx"
rgb = lambda h: RGBColor.from_string(h)
NAVY, SUB, MUTE = rgb("1F2A44"), rgb("4D5869"), rgb("7A8494")
BLUE, TEAL, LAV, CLAY = rgb("4A6A9C"), rgb("3E8C7E"), rgb("8A7DB0"), rgb("9A7B57")
VAL = rgb("A5524A")                     # muted brick, used only for the self-validation loop
PALE_BLUE, PALE_GREEN, PALE_LAV, GREY_BG, LINE = rgb("F3F6FA"), rgb("EEF5F1"), rgb("F5F3FA"), rgb("F6F7F9"), rgb("C9CFD8")
WHITE = rgb("FFFFFF")
FONT = "Arial"

prs = Presentation(HERE / "template" / "SIH2026-IDEA-template.pptx")
ids = prs.slides._sldIdLst
for i, sid in reversed(list(enumerate(list(ids)))):
    if i != 2: prs.part.drop_rel(sid.rId); ids.remove(sid)
s = prs.slides[0]
for sh in list(s.shapes):
    if sh.has_text_frame and "Technologies to be used" in sh.text_frame.text:
        sh._element.getparent().remove(sh._element)
    elif sh.has_text_frame and sh.text_frame.text.strip() == "TECHNICAL APPROACH":
        p = sh.text_frame.paragraphs[0]
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
        p.runs[0].text = "AquaSDR · TECHNICAL APPROACH"; p.runs[0].font.size = Pt(30)

I = Inches
def flat(shape):
    """No theme shadow/effects: remove <p:style> and set an explicit empty effect list."""
    el = shape._element
    st = el.find(qn("p:style"))
    if st is not None: el.remove(st)
    spPr = el.find(qn("p:spPr"))
    if spPr is not None and spPr.find(qn("a:effectLst")) is None:
        spPr.append(spPr.makeelement(qn("a:effectLst"), {}))
    return shape

def text(x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, size=9, color=NAVY, bold=False, italic=False, spacing=None):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, side, 0)
    if isinstance(paras, str): paras = [paras]
    for k, para in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); p.alignment = align
        if spacing: p.line_spacing = spacing
        for t, st in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run(); r.text = t; f = r.font; f.name = FONT
            f.size = Pt(st.get("size", size)); f.bold = st.get("bold", bold); f.italic = st.get("italic", italic); f.color.rgb = st.get("color", color)
    return tb

def box(x, y, w, h, fill=None, line=BLUE, lw=1.0, dashed=False, radius=0.05):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, I(x), I(y), I(w), I(h)); b.adjustments[0] = radius
    if fill is None: b.fill.background()
    else: b.fill.solid(); b.fill.fore_color.rgb = fill
    if line is None: b.line.fill.background()
    else:
        b.line.color.rgb = line; b.line.width = Pt(lw)
        if dashed: b.line.dash_style = MSO_LINE.DASH
    return flat(b)

def seg(x1, y1, x2, y2, color=NAVY, lw=1.25, head=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = color; c.line.width = Pt(lw); flat(c)
    if head:
        ln = c.line._get_or_add_ln(); ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c

def path(pts, color=NAVY, lw=1.25):
    for k in range(len(pts) - 1): seg(*pts[k], *pts[k + 1], color=color, lw=lw, head=(k == len(pts) - 2))

def dot(x, y, r=0.04, fill=WHITE, line=TEAL):
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, I(x - r), I(y - r), I(2 * r), I(2 * r))
    d.fill.solid(); d.fill.fore_color.rgb = fill; d.line.color.rgb = line; d.line.width = Pt(1.0); flat(d)

def num(x, y, n, color):
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, I(x + 0.02), I(y + 0.02), I(0.22), I(0.22))
    d.fill.solid(); d.fill.fore_color.rgb = color; d.line.fill.background(); flat(d)
    tf = d.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = n
    r.font.name = FONT; r.font.size = Pt(6.8); r.font.bold = True; r.font.color.rgb = WHITE

def photo(file, x, y, w, h, line=NAVY, lw=1.0):
    p = s.shapes.add_picture(str(file), I(x), I(y), I(w), I(h)); p.line.color.rgb = line; p.line.width = Pt(lw)
    return flat(p)

def header(x, y, w, t, color, tag=None):
    text(x, y, w, 0.24, [[(t, {"bold": True, "size": 10.5, "color": color})] + ([("   " + tag, {"size": 7.8, "italic": True, "color": MUTE})] if tag else [])])

# ---------------- subtitle under the template title ----------------
text(2.0, 0.93, 9.3, 0.24, [[("Adaptive Software-Defined Sonar Transmitter for AUVs", {"size": 12, "italic": True, "color": SUB})]], align=PP_ALIGN.CENTER)
text(2.0, 1.17, 9.3, 0.2, [[("REAL SENSORS  →  ON-CHIP LOGIC  →  REAL PAYLOAD", {"size": 8, "bold": True, "color": MUTE})]], align=PP_ALIGN.CENTER)

# =========================== LEFT · SENSE ===========================
header(0.30, 1.42, 3.3, "SENSE · REAL ENVIRONMENT", TEAL)
PX, PY, PW = 1.45, 1.72, 1.95
PH = PW / 0.909
photo(IMG / "sensors.jpg", PX, PY, PW, PH, line=TEAL)
P = lambda u, v: (PX + u / 1400 * PW, PY + v / 1541 * PH)
for label, (u, v), ly in [("TDS / SALINITY PROXY", (514, 424), 1.98), ("HC-SR04 · AIR RANGE", (719, 617), 2.38),
                          ("TURBIDITY", (540, 822), 2.78), ("DS18B20 TEMPERATURE", (1079, 796), 3.18), ("PRESSURE / DEPTH", (1028, 1130), 3.58)]:
    px, py = P(u, v)
    text(0.10, ly - 0.085, 1.30, 0.18, [[(label, {"bold": True, "size": 7.0})]], align=PP_ALIGN.RIGHT)
    seg(1.43, ly, px, py, color=TEAL, lw=0.9); dot(px, py)
text(PX, PY + PH + 0.04, PW, 0.18, [[("ADC  ·  1-Wire  ·  I²C", {"bold": True, "size": 8.5, "color": TEAL})]], align=PP_ALIGN.CENTER)

text(0.30, 4.16, 3.3, 0.18, [[("ADDITIONAL SENSORS", {"bold": True, "size": 7.8, "color": NAVY})]])
box(0.30, 4.36, 1.24, 0.62, fill=WHITE, line=TEAL)
photo(IMG / "bmp390.jpg", 0.34, 4.40, 0.52, 0.54, line=LINE, lw=0.5)
text(0.90, 4.47, 0.62, 0.45, [[("BMP390", {"bold": True, "size": 7.6})], [("cabin pressure", {"size": 6.8, "color": SUB})], [("I²C", {"size": 6.8, "color": SUB})]])
box(1.60, 4.36, 0.92, 0.62, fill=WHITE, line=TEAL)
text(1.64, 4.44, 0.84, 0.5, [[("IMU + magnetometer", {"bold": True, "size": 7.4})], [("ISM330 · MMC5983", {"size": 6.6, "color": SUB})], [("I²C", {"size": 6.6, "color": SUB})]], align=PP_ALIGN.CENTER)
box(2.58, 4.36, 0.97, 0.62, fill=WHITE, line=MUTE, dashed=True)
text(2.61, 4.46, 0.91, 0.5, [[("INA219 ×2", {"bold": True, "size": 7.6})], [("MOUNTED · NOT", {"size": 6.4, "color": MUTE, "bold": True})], [("COMMISSIONED", {"size": 6.4, "color": MUTE, "bold": True})]], align=PP_ALIGN.CENTER)

text(0.30, 5.13, 3.3, 0.18, [[("ONE ADAPTATION POLICY · TWO INPUTS", {"bold": True, "size": 7.8, "color": NAVY})]])
for k, (a, b) in enumerate([("A0 SENSOR", "TDS_PROXY → ADC"), ("WEB COMMAND", "console → UART")]):
    bx = 0.30 + k * 1.55
    box(bx, 5.33, 1.45, 0.42, fill=PALE_GREEN, line=TEAL)
    text(bx, 5.37, 1.45, 0.36, [[(a, {"bold": True, "size": 8})], [(b, {"size": 7, "color": SUB})]], align=PP_ALIGN.CENTER)
for bx in (1.025, 2.575): seg(bx, 5.75, bx, 5.90, color=TEAL, lw=1.0)
seg(1.025, 5.90, 3.66, 5.90, color=TEAL, lw=1.0)

# =========================== CENTRE · STM32 ===========================
header(3.75, 1.42, 4.7, "STM32F446RE · EMBEDDED CONTROL", BLUE, tag="methodology & process")
box(3.75, 1.72, 4.70, 4.62, fill=PALE_BLUE, line=BLUE, lw=1.0, radius=0.025)
box(3.85, 1.80, 4.50, 0.30, fill=NAVY, line=None, radius=0.2)
text(3.85, 1.80, 4.50, 0.30, [[("THE NEXT PING IS COMPUTED ON THE MCU", {"bold": True, "size": 11, "color": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
BX, BY, BW = 3.85, 2.18, 3.05
BH = BW / 2.7027
photo(IMG / "bay.jpg", BX, BY, BW, BH, line=BLUE)
for label, (fx, fy), (lx, ly) in [("STM32 NUCLEO-F446RE", (0.278, 0.114), (0.03, 0.70)), ("Sensor interface boards", (0.683, 0.82), (0.58, 0.08))]:
    tx, ty = BX + fx * BW, BY + fy * BH
    w = 0.058 * len(label) + 0.14; lx_, ly_ = BX + lx * BW, BY + ly * BH
    seg(lx_ + w / 2, ly_ + (0 if ly_ > ty else 0.17), tx, ty, color=WHITE, lw=1.0); dot(tx, ty, line=BLUE)
    box(lx_, ly_, w, 0.17, fill=WHITE, line=BLUE, lw=0.6, radius=0.5)
    text(lx_, ly_ + 0.005, w, 0.16, [[(label, {"bold": True, "size": 6.6})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
# evidence card
box(7.00, 2.18, 1.35, BH, fill=WHITE, line=LINE, lw=0.75, radius=0.06)
text(7.07, 2.20, 1.24, BH - 0.04, [[("MCU ADC CAPTURE", {"bold": True, "size": 7.8, "color": VAL})],
                                   [("LFM example · on-board", {"size": 6.8, "italic": True, "color": MUTE})],
                                   [("Centre  ", {"size": 7.4, "color": SUB}), ("42 kHz", {"bold": True, "size": 7.4})],
                                   [("Sweep  ", {"size": 7.4, "color": SUB}), ("4 kHz", {"bold": True, "size": 7.4})],
                                   [("Pulse  ", {"size": 7.4, "color": SUB}), ("2 ms", {"bold": True, "size": 7.4})],
                                   [("Window  ", {"size": 7.4, "color": SUB}), ("Hann", {"bold": True, "size": 7.4})],
                                   [("330/330", {"bold": True, "size": 7.2}), (" cmds ACKed", {"size": 7.0, "color": SUB})],
                                   [("0", {"bold": True, "size": 7.2}), (" DMA underruns", {"size": 7.0, "color": SUB})]], anchor=MSO_ANCHOR.MIDDLE)
steps = [("01", "ACQUIRE", "ADC · 1-Wire · I²C", "stale / out-of-range inputs rejected", BLUE),
         ("02", "DECIDE", "smoothing filter + hysteresis", "CLEAR / TRANSITION / MURKY", BLUE),
         ("03", "SYNTHESISE", "LFM chirp · geometric sweep · Barker-13", "Hann · Hamming · Blackman windows", BLUE),
         ("04", "STREAM", "TIM6 + DMA → 12-bit DAC1 / PA4", "1 MS/s · 0 underruns · CPU free between transfers", BLUE),
         ("06", "SELF-VALIDATE", "ADC + DMA capture of the DAC output", "FFT · spectrogram · matched filtering", VAL)]
Y0, DY, CX = 3.48, 0.56, 3.98
seg(CX + 0.13, Y0 + 0.13, CX + 0.13, Y0 + 3 * DY + 0.13, color=BLUE, lw=1.25)
for k, (n, t, d1, d2, col) in enumerate(steps):
    y = Y0 + k * DY
    num(CX, y, n, col)
    text(4.34, y - 0.02, 4.0, 0.18, [[(t, {"bold": True, "size": 9.2})]])
    text(4.34, y + 0.15, 4.0, 0.34, [[(d1, {"size": 7.6, "color": SUB})], [(d2, {"size": 7.6, "color": SUB})]])
path([(3.90, Y0 + 3 * DY + 0.13), (3.90, Y0 + 4 * DY + 0.13), (CX, Y0 + 4 * DY + 0.13)], color=VAL, lw=1.0)   # DAC pin -> capture
seg(PX + PW, Y0 + 0.13, 3.75, Y0 + 0.13, color=TEAL, lw=1.25, head=True)                    # sensors -> 01
path([(3.66, 5.90), (3.66, Y0 + DY + 0.13), (3.75, Y0 + DY + 0.13)], color=TEAL, lw=1.25)    # both inputs -> 02

# =========================== RIGHT · NEXT PING → CONDITION → PAYLOAD ===========================
header(8.65, 1.42, 3.0, "ADAPTED NEXT PING", LAV)
LX, CW, TY = 8.72, 0.69, 1.72
for j, (c, colr) in enumerate([("CLEAR", rgb("5B84B8")), ("TRANSITION", LAV), ("MURKY", CLAY)]):
    cx = LX + 0.86 + j * CW
    box(cx + 0.02, TY, CW - 0.04, 0.21, fill=colr, line=None, radius=0.3)
    text(cx, TY + 0.005, CW, 0.2, [[(c, {"bold": True, "size": 6.8, "color": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for i, (n, vals) in enumerate([("Centre", ["42 kHz", "40 kHz", "38 kHz"]), ("Bandwidth", ["4 kHz", "2 kHz", "1 kHz"]),
                               ("Pulse", ["2 ms", "4 ms", "8 ms"]), ("Amplitude", ["35 %", "50 %", "62 %"])]):
    y = TY + 0.27 + i * 0.215
    text(LX, y, 0.86, 0.2, [[(n, {"bold": True, "size": 8.4})]])
    for j, v in enumerate(vals): text(LX + 0.86 + j * CW, y, CW, 0.2, [[(v, {"size": 8.4, "color": SUB})]], align=PP_ALIGN.CENTER)
    if i < 3: seg(LX, y + 0.2, LX + 0.86 + 3 * CW, y + 0.2, color=rgb("E3E7ED"), lw=0.5)
text(8.72, 2.92, 4.4, 0.34, [[("CLEAR → shorter / sharper ping", {"size": 7.6, "color": SUB})],
                             [("MURKY → lower band / longer pulse / higher output demand", {"size": 7.6, "color": SUB})]])
path([(8.45, Y0 + DY + 0.13), (8.55, Y0 + DY + 0.13), (8.55, 2.40), (8.68, 2.40)], color=LAV, lw=1.25)   # 02 -> next ping

box(8.65, 3.34, 4.45, 1.16, fill=WHITE, line=LAV, lw=1.0, dashed=True, radius=0.05)
num(8.73, 3.40, "05", LAV)
text(9.06, 3.42, 3.4, 0.22, [[("CONDITION · ANALOG FRONT END", {"bold": True, "size": 9})]])
photo(IMG / "afe_chain.png", 8.76, 3.72, 2.50, 2.50 / 3.645, line=WHITE, lw=0.25)
text(11.36, 3.70, 1.7, 0.5, [[("2nd-order low-pass filter", {"bold": True, "size": 7.6})], [("70.7 kHz design target", {"size": 7.4, "color": SUB})], [("TLV9062 ×1.5", {"bold": True, "size": 7.6})]])
box(11.36, 4.20, 1.52, 0.2, fill=WHITE, line=LAV, lw=0.6, dashed=True, radius=0.5)
text(11.36, 4.205, 1.52, 0.19, [[("DESIGNED · NEXT BUILD", {"bold": True, "size": 6.6, "color": LAV})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
path([(8.45, Y0 + 3 * DY + 0.13), (8.58, Y0 + 3 * DY + 0.13), (8.58, 4.30), (8.65, 4.30)], color=LAV, lw=1.25)  # 04 DAC -> 05

seg(10.875, 4.50, 10.875, 4.64, color=NAVY, lw=1.25, head=True)
PW2 = 4.45; PH2 = PW2 / 3.735
photo(IMG / "payload.jpg", 8.65, 4.66, PW2, PH2, line=NAVY, lw=1.25)
text(8.65, 4.66 + PH2 + 0.04, PW2, 0.2, [[("AQUASDR PAYLOAD · PHYSICAL PROTOTYPE  ", {"bold": True, "size": 8.6}), ("acrylic payload hull", {"size": 7.6, "color": SUB})]])
text(8.65, 4.66 + PH2 + 0.24, PW2, 0.18, [[("External scope validation → PENDING   ·   AUV hull-slot pod → NEXT BUILD", {"size": 7.2, "italic": True, "color": MUTE})]])

# =========================== legend + bottom storyline ===========================
box(3.85, 6.405, 0.24, 0.12, fill=WHITE, line=NAVY, lw=0.8, radius=0.3); text(4.13, 6.385, 1.0, 0.16, [[("built / running", {"size": 6.8, "color": SUB})]])
box(5.10, 6.405, 0.24, 0.12, fill=WHITE, line=NAVY, lw=0.8, dashed=True, radius=0.3); text(5.38, 6.385, 1.3, 0.16, [[("designed / next build", {"size": 6.8, "color": SUB})]])
flow = [("SENSE", False), ("DECIDE", False), ("SYNTHESISE", False), ("STREAM", False), ("CONDITION", True), ("VERIFY", False)]
widths = [1.0, 1.0, 1.25, 1.0, 1.2, 1.0]; gap = 0.55
x = (13.333 - (sum(widths) + gap * (len(widths) - 1))) / 2
for k, ((t, dashed), w) in enumerate(zip(flow, widths)):
    box(x, 6.62, w, 0.25, fill=WHITE, line=NAVY if not dashed else LAV, lw=0.75, dashed=dashed, radius=0.12)
    text(x, 6.62, w, 0.25, [[(t, {"bold": True, "size": 8.4, "color": NAVY})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if k < len(flow) - 1: seg(x + w + 0.07, 6.74, x + w + gap - 0.07, 6.74, color=NAVY, lw=1.0, head=True)
    x += w + gap

prs.save(OUT)
print("wrote", OUT)
