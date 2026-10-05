#!/usr/bin/env python3
"""SIH 2026 deck, slide 3 (TECHNICAL APPROACH), v5: the complete system architecture as one story, no status labels
(implementation status and evidence live on slide 4). Derived from v4: same architecture as v2 in a restrained engineering-figure style.
Palette: navy + greyscale; teal only for sensor inputs; lavender only for designed / next-build. No shadows, no pills.
Layout: INPUT (left) -> EMBEDDED CONTROL (centre) -> PHYSICAL OUTPUT (right), TECHNOLOGIES TO BE USED band, minimal flow.
Only the team's own photographs; numbers from the repository's recorded evidence.
    python3 docs/deck/build_technical_slide_v3.py   -> docs/deck/AquaSDR_slide3_technical_approach_v5.pptx
"""
import os, pathlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn

HERE = pathlib.Path(__file__).parent
IMG = HERE / "img"
FIGURE_ONLY = os.environ.get("FIGURE_ONLY") == "1"   # just the architecture figure, no template decoration
OUT = HERE / ("AquaSDR_slide3_figure_v5.pptx" if FIGURE_ONLY else "AquaSDR_slide3_technical_approach_v5.pptx")
rgb = RGBColor.from_string
NAVY, TEAL, BG, LAV, CHAR, GRAY, WHITE = rgb("183153"), rgb("3B8C7A"), rgb("F3F6F9"), rgb("8A7BAF"), rgb("263238"), rgb("667085"), rgb("FFFFFF")
RULE = rgb("D0D5DD")
FONT = "Arial"

prs = Presentation(HERE / "template" / "SIH2026-IDEA-template.pptx")
ids = prs.slides._sldIdLst
for i, sid in reversed(list(enumerate(list(ids)))):
    if i != 2: prs.part.drop_rel(sid.rId); ids.remove(sid)
s = prs.slides[0]
for sh in list(s.shapes):
    if sh.has_text_frame and "Technologies to be used" in sh.text_frame.text:
        sh._element.getparent().remove(sh._element)
    elif sh.has_text_frame and sh.text_frame.text.strip() == "TECHNICAL APPROACH":
        p = sh.text_frame.paragraphs[0]
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
        p.runs[0].text = "AquaSDR · TECHNICAL APPROACH"; p.runs[0].font.size = Pt(30)

if FIGURE_ONLY:
    for sh in list(s.shapes): sh._element.getparent().remove(sh._element)
I = Inches
def flat(shape):
    el = shape._element
    st = el.find(qn("p:style"))
    if st is not None: el.remove(st)
    spPr = el.find(qn("p:spPr"))
    if spPr is not None and spPr.find(qn("a:effectLst")) is None: spPr.append(spPr.makeelement(qn("a:effectLst"), {}))
    return shape

def text(x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, size=8.5, color=CHAR, bold=False, italic=False):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, side, 0)
    if isinstance(paras, str): paras = [paras]
    for k, para in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); p.alignment = align
        for t, st in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run(); r.text = t; f = r.font; f.name = FONT
            f.size = Pt(st.get("size", size)); f.bold = st.get("bold", bold); f.italic = st.get("italic", italic); f.color.rgb = st.get("color", color)
    return tb

def rect(x, y, w, h, fill=None, line=NAVY, lw=0.75, dashed=False, radius=0.0):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
    if radius: b.adjustments[0] = radius
    if fill is None: b.fill.background()
    else: b.fill.solid(); b.fill.fore_color.rgb = fill
    if line is None: b.line.fill.background()
    else:
        b.line.color.rgb = line; b.line.width = Pt(lw)
        if dashed: b.line.dash_style = MSO_LINE.DASH
    return flat(b)

def seg(x1, y1, x2, y2, color=NAVY, lw=1.0, head=False, dashed=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = color; c.line.width = Pt(lw); flat(c)
    if dashed: c.line.dash_style = MSO_LINE.DASH
    if head:
        ln = c.line._get_or_add_ln(); ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "sm", "len": "sm"}))
    return c

def path(pts, color=NAVY, lw=1.0):
    for k in range(len(pts) - 1): seg(*pts[k], *pts[k + 1], color=color, lw=lw, head=(k == len(pts) - 2))

def dot(x, y, r=0.035, line=NAVY):
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, I(x - r), I(y - r), I(2 * r), I(2 * r))
    d.fill.solid(); d.fill.fore_color.rgb = WHITE; d.line.color.rgb = line; d.line.width = Pt(1.0); flat(d)

def photo(file, x, y, w, h, line=NAVY, lw=0.75):
    p = s.shapes.add_picture(str(file), I(x), I(y), I(w), I(h)); p.line.color.rgb = line; p.line.width = Pt(lw)
    return flat(p)

def tag(x, y, label="BUILT"):
    w = 0.075 * len(label) + 0.16
    rect(x, y, w, 0.17, fill=NAVY, line=None)
    text(x, y, w, 0.17, [[(label, {"bold": True, "size": 6.8, "color": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def header(x, y, w, label, subtitle):
    text(x, y, w, 0.18, [[(label, {"bold": True, "size": 9.6, "color": NAVY})]])
    text(x, y + 0.18, w, 0.17, [[(subtitle, {"size": 8.6, "color": GRAY})]])
    seg(x, y + 0.38, x + w, y + 0.38, color=RULE, lw=0.75)


(lambda *a, **k: None if FIGURE_ONLY else text(*a, **k))(2.0, 0.93, 9.3, 0.24, [[("Environment  →  Embedded Decision  →  Adaptive Ping  →  Physical Output", {"size": 12, "italic": True, "color": GRAY})]], align=PP_ALIGN.CENTER)

# =========================== SENSE ===========================
header(0.30, 1.30, 3.35, "SENSE · REAL ENVIRONMENT", "Live water-condition inputs")
PX, PY, PW = 1.35, 1.80, 2.30
PH = PW / 0.909
photo(IMG / "sensors.jpg", PX, PY, PW, PH)
P = lambda u, v: (PX + u / 1400 * PW, PY + v / 1541 * PH)
for label, (u, v), ly in [("TURBIDITY", (540, 822), 2.15), ("TDS / SALINITY PROXY", (514, 424), 2.65),
                          ("PRESSURE / DEPTH", (1028, 1130), 3.15), ("DS18B20 TEMPERATURE", (1079, 796), 3.65)]:
    px, py = P(u, v)
    ex = PX; ey = ly + (ex - 1.31) / (px - 1.31) * (py - ly)
    text(0.02, ly - 0.09, 1.27, 0.19, [[(label, {"bold": True, "size": 7.3, "color": NAVY})]], align=PP_ALIGN.RIGHT)
    seg(1.31, ly, ex, ey, color=NAVY, lw=0.9)
    seg(ex, ey, px, py, color=WHITE, lw=1.4); dot(px, py, r=0.042)
text(PX, PY + PH + 0.05, PW, 0.2, [[("ADC · 1-Wire · I²C", {"bold": True, "size": 9, "color": TEAL})]], align=PP_ALIGN.CENTER)
text(0.30, 4.72, 3.35, 0.5, [[("Environmental inputs provide the state used to select the next ping.", {"size": 9.2, "italic": True, "color": NAVY})]])

# =========================== STM32 (dominant) ===========================
header(3.80, 1.30, 4.60, "STM32F446RE · EMBEDDED CONTROL", "Firmware in C / C++ on the payload's own MCU")
rect(3.80, 1.76, 4.60, 3.98, fill=BG, line=NAVY, lw=1.0)
rect(3.80, 1.76, 4.60, 0.30, fill=NAVY, line=None)
text(3.80, 1.76, 4.60, 0.30, [[("THE NEXT PING IS COMPUTED FROM THE CURRENT ENVIRONMENT", {"bold": True, "size": 9.2, "color": WHITE})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
BW = 3.30; BH = BW / 2.7027; BX = 3.80 + (4.60 - BW) / 2; BY = 2.13
photo(IMG / "bay.jpg", BX, BY, BW, BH)
for label, (fx, fy), (lx, ly) in [("STM32 NUCLEO-F446RE", (0.278, 0.114), (0.03, 0.74)), ("Sensor interface boards", (0.683, 0.82), (0.45, 0.06))]:
    tx, ty = BX + fx * BW, BY + fy * BH
    w = 0.062 * len(label) + 0.14; lx_, ly_ = BX + lx * BW, BY + ly * BH
    seg(lx_ + w / 2, ly_ + (0 if ly_ > ty else 0.18), tx, ty, color=WHITE, lw=1.0); dot(tx, ty)
    rect(lx_, ly_, w, 0.18, fill=WHITE, line=NAVY, lw=0.5)
    text(lx_, ly_ + 0.01, w, 0.16, [[(label, {"bold": True, "size": 7.0, "color": NAVY})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
steps = [("01", "ACQUIRE", "Read environmental inputs · reject stale / invalid readings"),
         ("02", "DECIDE", "Smoothing + hysteresis → CLEAR / TRANSITION / MURKY"),
         ("03", "ADAPT", "Centre / bandwidth · pulse duration · amplitude"),
         ("04", "SYNTHESISE", "LFM · geometric · Barker-13  ×  Hann / Hamming / Blackman"),
         ("05", "STREAM", "TIM6 + DMA → 12-bit DAC1 / PA4 · 1 MS/s"),
         ("06", "VERIFY", "ADC + DMA capture → FFT · spectrogram · matched filter")]
Y0, DY, NX = 3.45, 0.375, 3.95
seg(NX + 0.12, Y0 + 0.21, NX + 0.12, Y0 + 5 * DY - 0.02, color=RULE, lw=1.0)
for k, (n, t, d) in enumerate(steps):
    y = Y0 + k * DY
    rect(NX - 0.02, y - 0.01, 0.28, 0.21, fill=BG, line=None)
    text(NX, y, 0.26, 0.2, [[(n, {"bold": True, "size": 9.2, "color": NAVY})]], align=PP_ALIGN.CENTER)
    text(4.32, y - 0.01, 1.25, 0.2, [[(t, {"bold": True, "size": 9.4, "color": NAVY})]])
    text(4.32, y + 0.17, 4.05, 0.18, [[(d, {"size": 7.9, "color": CHAR})]])
seg(PX + PW, Y0 + 0.09, 3.80, Y0 + 0.09, color=TEAL, lw=1.2, head=True)                                       # sensors -> ACQUIRE

# =========================== ADAPTIVE OUTPUT ===========================
header(8.65, 1.30, 2.95, "ADAPTIVE NEXT PING", "Parameters chosen per water condition")
LX, TW, TY, LW, CW, HH, RH = 8.65, 4.45, 1.84, 1.15, 1.10, 0.23, 0.20     # full-width bordered table
TH = HH + 4 * RH
rect(LX, TY, TW, TH, fill=WHITE, line=NAVY, lw=0.9)
rect(LX, TY, TW, HH, fill=BG, line=None)
seg(LX, TY + HH, LX + TW, TY + HH, color=NAVY, lw=0.75)
for j in range(4): seg(LX + LW + j * CW, TY, LX + LW + j * CW, TY + TH, color=RULE, lw=0.6) if j else seg(LX + LW, TY, LX + LW, TY + TH, color=NAVY, lw=0.6)
text(LX + 0.08, TY, LW - 0.1, HH, [[("PARAMETER", {"bold": True, "size": 7.6, "color": GRAY})]], anchor=MSO_ANCHOR.MIDDLE)
for j, c in enumerate(["CLEAR", "TRANSITION", "MURKY"]):
    text(LX + LW + j * CW, TY, CW, HH, [[(c, {"bold": True, "size": 8.0, "color": NAVY})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for i, (n, vals) in enumerate([("Centre", ["42 kHz", "40 kHz", "38 kHz"]), ("Bandwidth", ["4 kHz", "2 kHz", "1 kHz"]),
                               ("Pulse", ["2 ms", "4 ms", "8 ms"]), ("Amplitude", ["35 %", "50 %", "62 %"])]):
    y = TY + HH + i * RH
    text(LX + 0.08, y, LW - 0.1, RH, [[(n.upper(), {"bold": True, "size": 7.8, "color": NAVY})]], anchor=MSO_ANCHOR.MIDDLE)
    for j, v in enumerate(vals): text(LX + LW + j * CW, y, CW, RH, [[(v, {"size": 8.8, "color": CHAR})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < 3: seg(LX, y + RH, LX + TW, y + RH, color=RULE, lw=0.5)
text(LX, TY + TH + 0.05, 4.45, 0.34, [[("CLEAR → shorter / sharper ping", {"size": 8.0, "color": CHAR})], [("MURKY → lower band / longer pulse / higher output demand", {"size": 8.0, "color": CHAR})]])
path([(8.40, Y0 + 2 * DY + 0.09), (8.47, Y0 + 2 * DY + 0.09), (8.47, TY + TH / 2), (8.645, TY + TH / 2)], color=NAVY, lw=1.0)      # ADAPT -> parameters

rect(8.65, 3.32, 4.45, 0.86, fill=WHITE, line=NAVY, lw=0.75)
text(8.73, 3.37, 3.6, 0.18, [[("CONDITION · ANALOG FRONT END", {"bold": True, "size": 9.0, "color": NAVY})]])
photo(IMG / "afe_chain.png", 8.73, 3.59, 2.05, 2.05 / 3.645, line=WHITE, lw=0.25)
text(10.90, 3.60, 2.15, 0.55, [[("LPF → TLV9062 → output", {"bold": True, "size": 8.4, "color": NAVY})], [("2nd-order low-pass · 70.7 kHz", {"size": 7.8, "color": GRAY})],
                               [("op-amp gain ×1.5", {"size": 7.8, "color": GRAY})]])
path([(8.40, Y0 + 4 * DY + 0.09), (8.58, Y0 + 4 * DY + 0.09), (8.58, 3.75), (8.65, 3.75)], color=NAVY, lw=1.0)    # STREAM -> analog

seg(10.875, 4.18, 10.875, 4.28, color=NAVY, lw=1.0, head=True)
PW2 = 3.70; PH2 = PW2 / 3.735
photo(IMG / "payload.jpg", 8.65 + (4.45 - PW2) / 2, 4.29, PW2, PH2, lw=1.0)
text(8.65, 4.29 + PH2 + 0.03, 4.45, 0.18, [[("AQUASDR PAYLOAD", {"bold": True, "size": 8.8, "color": NAVY}), ("   physical acrylic payload prototype for AUV integration", {"size": 7.8, "color": GRAY})]])
VY = 4.29 + PH2 + 0.23
seg(10.875, VY - 0.03, 10.875, VY + 0.04, color=NAVY, lw=1.0, head=True)
rect(8.65, VY + 0.05, 4.45, 0.36, fill=BG, line=NAVY, lw=0.75)
text(8.73, VY + 0.06, 4.33, 0.34, [[("VALIDATION  ", {"bold": True, "size": 8.4, "color": NAVY}), ("MCU ADC capture · FFT · spectrogram · matched filter", {"size": 7.8, "color": CHAR})],
                                   [("+ external instrument (oscilloscope / spectrum analyzer)", {"size": 7.2, "color": GRAY})]], anchor=MSO_ANCHOR.MIDDLE)

# =========================== TECHNOLOGIES ===========================
TY2 = 5.98
rect(0.30, TY2, 12.80, 0.50, fill=WHITE, line=RULE, lw=0.75)
text(0.40, TY2 + 0.07, 1.30, 0.4, [[("TECHNOLOGIES", {"bold": True, "size": 8.6, "color": NAVY})], [("TO BE USED", {"bold": True, "size": 8.6, "color": NAVY})]])
seg(1.75, TY2 + 0.07, 1.75, TY2 + 0.45, color=RULE, lw=0.75)
techs = [("Embedded", "C/C++ · STM32F446RE · ADC · TIM6 · DMA · 12-bit DAC"), ("Signal processing", "LFM · Geometric sweep · Barker-13 · FFT · Windowing"),
         ("Analog", "Sallen-Key LPF · TLV9062"), ("Software", "React · Node.js · Python · OpenSCAD")]
for k, (a_, b_) in enumerate(techs):
    cx, cy = 1.90 + (k % 2) * 5.6, TY2 + 0.05 + (k // 2) * 0.23
    text(cx, cy, 5.5, 0.21, [[(a_ + "   ", {"bold": True, "size": 8.4, "color": NAVY}), (b_, {"size": 8.4, "color": CHAR})]])

# =========================== bottom story ===========================
flow = ["SENSE", "DECIDE", "ADAPT", "SYNTHESISE", "STREAM", "CONDITION", "VERIFY"]
widths = [0.86, 0.90, 0.80, 1.18, 0.95, 1.12, 0.90]; gap = 0.50
x = (13.333 - (sum(widths) + gap * (len(widths) - 1))) / 2
for k, (t, w) in enumerate(zip(flow, widths)):
    rect(x, 6.58, w, 0.26, fill=WHITE, line=NAVY, lw=0.75)
    text(x, 6.58, w, 0.26, [[(t, {"bold": True, "size": 8.6, "color": NAVY})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if k < len(flow) - 1: seg(x + w + 0.07, 6.71, x + w + gap - 0.07, 6.71, color=NAVY, lw=0.9, head=True)
    x += w + gap

prs.save(OUT)
print("wrote", OUT)
